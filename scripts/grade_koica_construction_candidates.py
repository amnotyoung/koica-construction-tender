#!/usr/bin/env python3
"""Grade 127 KOICA construction candidates using auditable attachment evidence.

The grading boundary is deliberately conservative:

* A: line-item quantities and positive line-item rates/amounts are recoverable.
* B: a machine-readable BOQ with line-item quantities is recoverable, but usable
  rates are absent or not verified.
* C: only scope, area, ceiling, document-trace, or no linked attachment exists.

Grades describe project-specific estimating readiness.  They are not claims
about contract/final cost, country-wide price representativeness, or semantic
use of every attachment.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

warnings.filterwarnings(
    "ignore",
    message=r"Print area cannot be set to Defined name:.*",
    category=UserWarning,
)

try:
    import xlrd
except ImportError:  # bundled Python may not include legacy XLS support
    xlrd = None


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CANDIDATES = Path(
    "outputs/koica-official-open-data/KOICA_공식데이터_건축관련후보.csv"
)
DEFAULT_OUTPUT_DIR = Path("outputs/koica-candidate-grading")

SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
BOQ_FILE_RE = re.compile(
    r"(?i)(\bboq\b|bill.{0,8}(?:quantit|price)|schedule.{0,8}(?:price|quantit)|"
    r"priced|unpriced|blank.{0,5}bill|내역서|물량|수량|견적|단가|공사비)"
)
BLANK_FILE_RE = re.compile(r"(?i)(blank|unpriced|not.?priced|공내역)")
CONSTRUCTION_TITLE_RE = re.compile(
    r"(?i)(construction|civil.?works|renovation|공사|시공|신축|개보수|리모델링)"
)
NON_CONSTRUCTION_TITLE_RE = re.compile(
    r"(?i)(design|supervision|consult|supply|procurement|설계|감리|PMC|"
    r"기자재|물품|가구|차량|구매|공급)"
)
TOTAL_RE = re.compile(
    r"(?i)^(?:sub\s*total|subtotal|grand\s*total|total|합계|소계|총계|"
    r"итого|total general|total parcial)\b"
)

DESCRIPTION_RE = re.compile(
    r"(?i)(description|designation|item.?description|description.?of.?work|"
    r"name.?of.?work|scope.?of.?work|품명|항목|내역|공종|명칭|"
    r"описание|наименование|descripci[oó]n|concepto|detalle)"
)
QUANTITY_RE = re.compile(
    r"(?i)(^|\b)(qty\.?|quantity|quantities|수량|количеств|cantidad|quantit[eé])($|\b)"
)
UNIT_RE = re.compile(
    r"(?i)(^|\b)(unit|uom|단위|ед\.?|единиц|unidad|unit[eé])($|\b)"
)
RATE_RE = re.compile(
    r"(?i)(unit.?rate|unit.?price|rate|단가|재료비|노무비|경비|цена|"
    r"precio.?unitario|prix.?unitaire)"
)
AMOUNT_RE = re.compile(
    r"(?i)(amount|total.?price|extended.?price|금액|합계금액|стоимост|"
    r"importe|monto|montant|total.?amount)"
)


@dataclass
class TableSignal:
    sheet: str
    sheet_state: str
    header_row: int
    quantity_rows: int
    rate_rows: int
    amount_rows: int
    priced_rows: int
    priced_ratio: float
    description_columns: str
    quantity_columns: str
    rate_columns: str
    amount_columns: str


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        if re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
            try:
                return float(text)
            except ValueError:
                return None
    return None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def excel_column(index: int) -> str:
    index += 1
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def joined_header(matrix: list[list[Any]], start: int, column: int) -> str:
    parts = []
    for row_no in range(start, min(start + 3, len(matrix))):
        if column < len(matrix[row_no]):
            value = clean(matrix[row_no][column])
            if value:
                parts.append(value)
    return " ".join(parts)


def find_table_signals(
    matrix: list[list[Any]], sheet: str, sheet_state: str = "visible"
) -> list[TableSignal]:
    """Find BOQ-like tables in a worksheet matrix.

    This function is intentionally exposed for synthetic tests.  It requires
    positive numeric quantities and descriptive text below a recognized header.
    """

    if not matrix:
        return []
    max_columns = min(max((len(row) for row in matrix), default=0), 80)
    signals: list[TableSignal] = []
    last_end = -1
    for header_index in range(min(len(matrix), 250)):
        if header_index <= last_end:
            continue
        headers = [joined_header(matrix, header_index, col) for col in range(max_columns)]
        description_cols = [i for i, value in enumerate(headers) if DESCRIPTION_RE.search(value)]
        quantity_cols = [i for i, value in enumerate(headers) if QUANTITY_RE.search(value)]
        unit_cols = [i for i, value in enumerate(headers) if UNIT_RE.search(value)]
        rate_cols = [i for i, value in enumerate(headers) if RATE_RE.search(value)]
        amount_cols = [i for i, value in enumerate(headers) if AMOUNT_RE.search(value)]
        if not quantity_cols or not (description_cols or unit_cols):
            continue
        if not (rate_cols or amount_cols or description_cols):
            continue

        quantity_rows = 0
        rate_rows = 0
        amount_rows = 0
        priced_rows = 0
        empty_run = 0
        scan_start = header_index + 1
        scan_end = min(len(matrix), scan_start + 2000)
        actual_end = scan_start
        for row_index in range(scan_start, scan_end):
            row = matrix[row_index]
            quantity_values = [
                number(row[col]) if col < len(row) else None for col in quantity_cols
            ]
            positive_quantity = any(value is not None and value > 0 for value in quantity_values)
            description = " ".join(
                clean(row[col]) for col in description_cols if col < len(row)
            )
            if not description:
                left_edge = min(quantity_cols)
                description = " ".join(clean(value) for value in row[:left_edge] if clean(value))
            descriptive = bool(description) and not TOTAL_RE.search(description)
            if positive_quantity and descriptive:
                quantity_rows += 1
                rate_positive = any(
                    (value := number(row[col] if col < len(row) else None)) is not None
                    and value > 0
                    for col in rate_cols
                )
                amount_positive = any(
                    (value := number(row[col] if col < len(row) else None)) is not None
                    and value > 0
                    for col in amount_cols
                )
                rate_rows += int(rate_positive)
                amount_rows += int(amount_positive)
                priced_rows += int(rate_positive or amount_positive)
                empty_run = 0
                actual_end = row_index
            elif any(clean(value) for value in row):
                empty_run = 0
            else:
                empty_run += 1
                if empty_run >= 25 and quantity_rows >= 3:
                    break
        if quantity_rows < 3:
            continue
        priced_ratio = priced_rows / quantity_rows if quantity_rows else 0.0
        signals.append(
            TableSignal(
                sheet=sheet,
                sheet_state=sheet_state,
                header_row=header_index + 1,
                quantity_rows=quantity_rows,
                rate_rows=rate_rows,
                amount_rows=amount_rows,
                priced_rows=priced_rows,
                priced_ratio=priced_ratio,
                description_columns=";".join(excel_column(i) for i in description_cols),
                quantity_columns=";".join(excel_column(i) for i in quantity_cols),
                rate_columns=";".join(excel_column(i) for i in rate_cols),
                amount_columns=";".join(excel_column(i) for i in amount_cols),
            )
        )
        last_end = actual_end
    return signals


def openpyxl_matrices(path: Path) -> Iterable[tuple[str, str, list[list[Any]]]]:
    workbook = load_workbook(
        path,
        data_only=True,
        read_only=True,
        keep_links=False,
    )
    try:
        for sheet in workbook.worksheets:
            max_row = min(sheet.max_row or 0, 5000)
            max_col = min(sheet.max_column or 0, 80)
            if max_row == 0 or max_col == 0:
                continue
            matrix = [
                list(row)
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=max_row,
                    min_col=1,
                    max_col=max_col,
                    values_only=True,
                )
            ]
            yield sheet.title, sheet.sheet_state, matrix
    finally:
        workbook.close()


def xlrd_matrices(path: Path) -> Iterable[tuple[str, str, list[list[Any]]]]:
    if xlrd is None:
        raise RuntimeError("xlrd is required to inspect legacy .xls attachments")
    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        for sheet in workbook.sheets():
            max_row = min(sheet.nrows, 5000)
            max_col = min(sheet.ncols, 80)
            if max_row == 0 or max_col == 0:
                continue
            matrix = [sheet.row_values(row, 0, max_col) for row in range(max_row)]
            visibility = getattr(sheet, "visibility", 0)
            state = "visible" if visibility == 0 else "hidden"
            yield sheet.name, state, matrix
    finally:
        workbook.release_resources()


def resolve_source_file(project_root: Path, source_file: str) -> Path | None:
    for base in (project_root / "data/unpacked", project_root / "data/raw"):
        path = base / source_file
        if path.exists() and path.is_file():
            return path
    return None


def scan_spreadsheet(path: Path) -> dict[str, Any]:
    signals: list[TableSignal] = []
    error = ""
    try:
        matrices = xlrd_matrices(path) if path.suffix.lower() == ".xls" else openpyxl_matrices(path)
        for sheet_name, sheet_state, matrix in matrices:
            signals.extend(find_table_signals(matrix, sheet_name, sheet_state))
    except Exception as exc:  # corrupt/encrypted legacy files remain traceable
        error = f"{type(exc).__name__}: {exc}"
    best = max(signals, key=lambda item: (item.quantity_rows, item.priced_rows), default=None)
    quantity_rows = max((item.quantity_rows for item in signals), default=0)
    priced_rows = max((item.priced_rows for item in signals), default=0)
    strong_priced = any(
        item.priced_rows >= 3 and item.priced_ratio >= 0.90 for item in signals
    )
    return {
        "scan_error": error,
        "table_count": len(signals),
        "quantity_rows_max": quantity_rows,
        "priced_rows_max": priced_rows,
        "strong_priced_table": int(strong_priced),
        "best_sheet": best.sheet if best else "",
        "best_sheet_state": best.sheet_state if best else "",
        "best_header_row": best.header_row if best else "",
        "best_priced_ratio": best.priced_ratio if best else 0.0,
        "table_signals": [asdict(item) for item in signals],
    }


def project_key(row: dict[str, str]) -> str:
    project_no = row.get("project_no_linked", "").strip()
    if project_no:
        return project_no
    alias = SOURCE_RECORD_PROJECT_ALIASES.get(row.get("source_record_id", ""))
    if alias:
        return alias
    return f"UNLINKED:{row['source_record_id']}"


def source_grade_hint(row: dict[str, str]) -> str:
    if row.get("unit_cost_potential") == "1":
        return "C1 면적·금액 원단위"
    if row.get("gross_floor_area_m2") and row.get("amount_value"):
        return "U1 면적·금액 범위불일치 가능"
    if row.get("amount_value"):
        return "C2 금액·범위"
    if row.get("existing_db_project_bid_nos"):
        return "C3 첨부 문서추적"
    return "C4 미연결·기초정보"


def construction_bid_relevant(detail: dict[str, Any] | None) -> bool:
    if not detail:
        return False
    contract_type = clean(detail.get("contract_type"))
    title = clean(detail.get("title"))
    if contract_type == "공사":
        return True
    return bool(CONSTRUCTION_TITLE_RE.search(title)) and not bool(
        NON_CONSTRUCTION_TITLE_RE.search(title)
    )


MANUAL_PROJECT_REVIEWS: dict[str, dict[str, Any]] = {
    "2021-00016": {
        "grade": "A",
        "technical_status": "A",
        "scope_status": "FULL",
        "price_stage": "DESIGN_ESTIMATE",
        "verification_level": "V3",
        "grade_detail": (
            "수동검증완료: 공사 입찰의 집행한도 예산서에서 주요 공종 BOQ 수량·GHS/USD 단가·금액 확인"
        ),
        "evidence_file": "L2025-00018-1/3. 집행한도금액 예산(안).xlsx",
        "evidence_location": "MAIN BUILDING!A3:H277 및 Grand Summary!A1:G18",
        "blocking_issue": "설계예산이며 낙찰·계약·준공가 아님; 가격시점·세금·환율 재확인",
        "next_action": "계약·준공 금액 확보 후 설계예산 대비율 산정",
        "verified_cost_amount": 4333426.767765368,
        "verified_cost_currency": "USD",
        "verified_cost_per_gfa_m2": 1666.70260298668,
        "verified_item_note": "주건물 208행; 6개 시설시트 자동검출 합계 320행(수동 품목 정제 필요)",
    },
    "2020-00119": {
        "grade": "A",
        "technical_status": "A-부분",
        "scope_status": "PARTIAL_BOUNDED",
        "price_stage": "DESIGN_ESTIMATE",
        "verification_level": "V3",
        "grade_detail": (
            "수동검증완료: 숨김 설계단가 644개와 건축 Bill No.01 BOQ 561행을 품목코드로 561/561 결합"
        ),
        "evidence_file": (
            "L2025-00085-1/.../BILL No1 -Architectural Works-01. HOSPITAL(BLANK BILL).xlsx"
        ),
        "evidence_location": "Works particular bill 및 숨김 일위대가목록·단가대비표",
        "blocking_issue": "건축 Bill No.01만 포함; 토목·기계·전기·통신·소방 등 제외",
        "next_action": "나머지 Bill을 동일 방식으로 복원하고 계약·준공가 확보",
        "verified_cost_amount": 8372712.422776788,
        "verified_cost_currency": "USD",
        "verified_cost_per_gfa_m2": 1526.196212682608,
        "verified_item_note": "건축 Bill No.01 561행·숨김단가 644개·561/561 결합",
    },
    "2013-00064": {
        "grade": "B",
        "technical_status": "B",
        "scope_status": "FULL",
        "price_stage": "NONE",
        "verification_level": "V3",
        "grade_detail": "수동검증완료: 건축·구조·전기·위생·조경 blank BOQ에 수량은 있으나 캐시 단가·금액은 0/공란",
        "evidence_file": "L2025-00012-1/.../Section 4. Q-Lab_Chacala_BOQ_blanked/",
        "evidence_location": "A1 ARCHITECTURE, A2/A3 STRUCTURE, B ELECTRIC, F HYDRO, G LANDSCAPE",
        "blocking_issue": "품목별 현지 단가 부재",
        "next_action": "볼리비아 현지 재료·노무·장비 단가 결합",
    },
    "2020-00070": {
        "grade": "B",
        "technical_status": "B-부분",
        "scope_status": "PARTIAL_BOUNDED",
        "price_stage": "NONE",
        "verification_level": "V3",
        "grade_detail": "수동검증완료: MVC 직업훈련센터 899행 blank BOQ에 수량은 있으나 단가·RD$/USD 금액은 공란",
        "evidence_file": "L2024-00078-1/Section 4_BOQ (BLANK sin PRICE)- MVC PROJECT_fin.xlsx",
        "evidence_location": "BOQ Code/Unit/Description/Quantity/Unit Cost/Total 표",
        "blocking_issue": "몬테크리스티 시설 부분자료이며 현지 단가 부재",
        "next_action": "도미니카 현지단가 결합 및 하이나 시설과 범위 분리",
    },
    "2020-00086": {
        "grade": "B",
        "technical_status": "B",
        "scope_status": "FULL",
        "price_stage": "NONE",
        "verification_level": "V3",
        "grade_detail": "수동검증완료: Urgench VTC·TTC Not-Priced 공내역서에 대규모 공종별 수량은 있으나 캐시 단가는 0/공란",
        "evidence_file": "L2023-00036-1/.../Section 4. Bill of Quantities(Not-Priced)/공내역서/",
        "evidence_location": "VTC·TTC 구조·건축·설비 공종별 XLS/XLSX",
        "blocking_issue": "품목별 현지 단가 부재; 재공고 파일 중복 제거 필요",
        "next_action": "고유 SHA 기준 수량 통합 후 우즈베키스탄 현지단가 결합",
    },
    "2023-00087": {
        "grade": "B",
        "technical_status": "B-부분",
        "scope_status": "PARTIAL_BOUNDED",
        "price_stage": "NONE",
        "verification_level": "V3",
        "grade_detail": "수동검증완료: 자이툰 도서관 2단계 리모델링 non-priced BOQ 48개 수량행, 단가 0",
        "evidence_file": "L2025-00017-2/1. BOQ for Zaytun Library_non-priced for 2 stages.xlsx",
        "evidence_location": "Sheet1 수량·단가 표",
        "blocking_issue": "도서관 부분자료이며 품목별 현지 단가 부재",
        "next_action": "이라크 단가 라이브러리와 품목·시점·지역 적합성 검토 후 결합",
    },
    "2022-00027": {
        "grade": "B",
        "technical_status": "B",
        "scope_status": "FULL",
        "price_stage": "NONE",
        "verification_level": "V3",
        "grade_detail": "수동검증완료: 국립소아병원 신축 Section 4 blank BOQ에 공종별 수량, 단가 0",
        "evidence_file": "L2025-00073-1/.../Section 4_BLANK BOQ.xlsx",
        "evidence_location": "Bill 1.0~6.6, 대표 Bill 4.3 수량표",
        "blocking_issue": "품목별 캄보디아 현지 단가 부재",
        "next_action": "공종별 BOQ 통합 후 현지단가 결합",
    },
    "2022-00076": {
        "grade": "B",
        "technical_status": "B",
        "scope_status": "FULL",
        "price_stage": "NONE",
        "verification_level": "V3",
        "grade_detail": "수동검증완료: 피지 국립재활센터 Section 4 BOQ에 공종별 수량, 단가 0",
        "evidence_file": "L2025-00051-1/Section 4 - Bill of Quantity.xlsx",
        "evidence_location": "공종별 시트, 대표 Mechanical Services 수량표",
        "blocking_issue": "품목별 피지 현지 단가 부재; 재공고 파일 해시 동일",
        "next_action": "고유 BOQ 1건으로 처리하고 피지 현지단가 결합",
    },
}

# Conservative title/country aliases verified against the local bid corpus.
# The source CSV intentionally leaves these project links blank.
SOURCE_RECORD_PROJECT_ALIASES = {
    "15085055:103": "2020-00070",
    "15085055:325": "2020-00070",
    "15085055:186": "2020-00070",
    "15085055:170": "2020-00070",
    "15085055:156": "2020-00070",
    "15085055:209": "2024-00053",
    "15085055:144": "2024-00053",
    "15073135:42": "2021-00085",
    "15073135:43": "2021-00085",
    "15085055:66": "2021-00085",
    "15085055:130": "2024-00006",
    "15085055:119": "2019-03655",
    "15085055:213": "2024-00107",
}


def record_cost_status(row: dict[str, str], project: dict[str, Any]) -> str:
    """Classify whether the official record itself is a construction-cost sample."""

    source = row.get("source_dataset", "")
    category = clean(row.get("procurement_category"))
    title_scope = f"{row.get('record_title', '')} {row.get('record_scope', '')}"
    if source == "aid_procurement_contracts":
        return "X-용역계약"
    if source == "country_project_reports":
        return "U-전체사업정보"
    if category == "물품":
        return "X-물품"
    if category == "용역" and NON_CONSTRUCTION_TITLE_RE.search(title_scope) and not re.search(
        r"(?i)(시공\s*용역|construction\s+works)", title_scope
    ):
        return "X-설계·감리·기타용역"
    if project["technical_status"] in {"A", "A-부분", "B", "B-부분"}:
        return project["technical_status"]
    if row.get("unit_cost_potential") == "1" and category == "공사":
        return "C?"
    return "U-근거부족"


def grade_project(
    key: str,
    candidate_rows: list[dict[str, str]],
    file_rows: list[dict[str, Any]],
    linked_bids: set[str] | None = None,
) -> dict[str, Any]:
    construction_file_rows = [
        row for row in file_rows if row.get("construction_bid_cost_relevant")
    ]
    unique_scans: dict[str, dict[str, Any]] = {}
    for file_row in construction_file_rows:
        digest = file_row.get("sha256", "") or file_row["source_file"]
        if digest not in unique_scans:
            unique_scans[digest] = file_row
    scanned = list(unique_scans.values())
    quantity_candidates = [row for row in scanned if row.get("quantity_rows_max", 0) >= 3]
    priced_candidates = [row for row in scanned if row.get("strong_priced_table")]
    boq_named = [
        row for row in construction_file_rows if row.get("boq_filename_candidate")
    ]
    scan_errors = [row for row in scanned if row.get("scan_error")]

    manual_status = "자동판정(V2)"
    price_stage = "NONE_UNKNOWN"
    scope_status = "UNKNOWN"
    verification_level = "V2" if construction_file_rows else "V0"
    technical_status = "U"
    evidence_file = ""
    evidence_location = ""
    blocking_issue = ""
    next_action = ""
    verified_cost_amount: float | None = None
    verified_cost_currency = ""
    verified_cost_per_gfa_m2: float | None = None
    verified_item_note = ""
    if key in MANUAL_PROJECT_REVIEWS:
        manual = MANUAL_PROJECT_REVIEWS[key]
        grade = manual["grade"]
        technical_status = manual["technical_status"]
        detail = manual["grade_detail"]
        manual_status = "수동검증완료"
        price_stage = manual["price_stage"]
        scope_status = manual["scope_status"]
        verification_level = manual["verification_level"]
        evidence_file = manual["evidence_file"]
        evidence_location = manual["evidence_location"]
        blocking_issue = manual["blocking_issue"]
        next_action = manual["next_action"]
        verified_cost_amount = manual.get("verified_cost_amount")
        verified_cost_currency = manual.get("verified_cost_currency", "")
        verified_cost_per_gfa_m2 = manual.get("verified_cost_per_gfa_m2")
        verified_item_note = manual.get("verified_item_note", "")
    elif priced_candidates:
        grade = "A"
        best = max(priced_candidates, key=lambda row: row.get("priced_rows_max", 0))
        technical_status = "A?"
        detail = (
            "자동 A후보: 수량행과 양수 단가/금액행이 함께 검출됨; "
            f"원문표 수동확인 필요({best.get('best_sheet', '')})"
        )
        manual_status = "A후보 수동검토필요"
        price_stage = "UNKNOWN"
        evidence_file = best.get("source_file", "")
        evidence_location = f"{best.get('best_sheet', '')}!header row {best.get('best_header_row', '')}"
        blocking_issue = "가격단계·통화·공사범위 수동확인 필요"
        next_action = "원문을 열어 최신 실제 BOQ와 가격 범위를 확인"
    elif quantity_candidates:
        grade = "B"
        best = max(quantity_candidates, key=lambda row: row.get("quantity_rows_max", 0))
        technical_status = "B?"
        detail = (
            "기계판독 BOQ 수량행 확인, 양수 단가/금액 결합은 미확인; "
            f"현지단가 입력 필요({best.get('best_sheet', '')})"
        )
        price_stage = "NONE_UNKNOWN"
        evidence_file = best.get("source_file", "")
        evidence_location = f"{best.get('best_sheet', '')}!header row {best.get('best_header_row', '')}"
        blocking_issue = "단가 부재 및 전체 공사범위 여부 수동확인 필요"
        next_action = "BOQ 범위를 확인하고 현지 또는 타 KOICA 단가를 결합"
    else:
        grade = "C"
        hint = min((source_grade_hint(row) for row in candidate_rows), default="C4 미연결·기초정보")
        if hint.startswith("C1"):
            technical_status = "C?"
            price_stage = "PLAN_CEILING"
            verification_level = "V1"
        else:
            technical_status = "U"
        detail = f"{hint}: 직접 수량×단가 산정 근거 미확인"
        if boq_named:
            detail += "; BOQ 명칭 파일은 있으나 구조화 수량표 자동확인 실패"
        blocking_issue = "품목 수량·단가 또는 범위 일치 GFA·공사금액 부족"
        next_action = "첨부 추가수집·OCR·수동검토 또는 설계/계약 BOQ 확보"

    country = next((row.get("country_ko", "") for row in candidate_rows if row.get("country_ko")), "미상")
    construction_titles = [
        row.get("record_title", "")
        for row in candidate_rows
        if row.get("procurement_category") == "공사"
    ]
    title = max(
        construction_titles
        or [row.get("record_title", "") for row in candidate_rows],
        key=len,
        default="",
    )
    representative_bid_title = max(
        {
            row.get("bid_title", "")
            for row in construction_file_rows
            if row.get("bid_title")
        },
        key=len,
        default="",
    )
    bids = sorted(
        {
            bid
            for row in candidate_rows
            for bid in row.get("existing_db_project_bid_nos", "").split(";")
            if bid
        }
        | {row["bid_no"] for row in file_rows if row.get("bid_no")}
        | (linked_bids or set())
    )
    areas = sorted(
        {
            float(row["gross_floor_area_m2"])
            for row in candidate_rows
            if number(row.get("gross_floor_area_m2")) is not None
        }
    )
    return {
        "project_key": key,
        "project_no": "" if key.startswith("UNLINKED:") else key,
        "country_ko": country,
        "representative_title": title,
        "representative_construction_bid_title": representative_bid_title,
        "grade": grade,
        "technical_status": technical_status,
        "grade_detail": detail,
        "price_stage": price_stage,
        "scope_status": scope_status,
        "verification_level": verification_level,
        "review_status": manual_status,
        "evidence_file": evidence_file,
        "evidence_location": evidence_location,
        "blocking_issue": blocking_issue,
        "next_action": next_action,
        "verified_cost_amount": verified_cost_amount,
        "verified_cost_currency": verified_cost_currency,
        "verified_cost_per_gfa_m2": verified_cost_per_gfa_m2,
        "verified_item_note": verified_item_note,
        "candidate_record_count": len(candidate_rows),
        "linked_bid_count": len(bids),
        "linked_bid_nos": ";".join(bids),
        "indexed_attachment_count": len(file_rows),
        "construction_attachment_count": len(construction_file_rows),
        "unique_attachment_count": len({row.get('sha256') or row['source_file'] for row in file_rows}),
        "spreadsheet_count": sum(row.get("extension") in SPREADSHEET_EXTENSIONS for row in construction_file_rows),
        "boq_named_file_count": len(boq_named),
        "quantity_table_file_count": len(quantity_candidates),
        "priced_table_file_count": len(priced_candidates),
        "max_quantity_rows_detected": max((row.get("quantity_rows_max", 0) for row in scanned), default=0),
        "max_priced_rows_detected": max((row.get("priced_rows_max", 0) for row in scanned), default=0),
        "scan_error_file_count": len(scan_errors),
        "gross_floor_area_m2_values": ";".join(f"{value:g}" for value in areas),
        "direct_project_estimate_ready": int(technical_status in {"A", "A-부분"}),
        "needs_local_rates": int(grade == "B"),
        "only_screening_or_trace": int(grade == "C"),
        "sample_weight": 1 if verification_level == "V3" and grade in {"A", "B", "C"} else 0,
        "claim_boundary": (
            "A도 첨부 가격단계 범위의 사업별 추정만 가능; 국가단가·계약가·준공가를 의미하지 않음"
        ),
    }


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def run(project_root: Path, candidates_path: Path, output_dir: Path) -> dict[str, Any]:
    with candidates_path.open(encoding="utf-8-sig", newline="") as handle:
        candidates = list(csv.DictReader(handle))
    if len(candidates) != 127:
        raise RuntimeError(f"expected 127 candidates, found {len(candidates)}")

    manifests = project_root / "data/manifests"
    file_index = json.loads((manifests / "file_index.json").read_text())
    details = json.loads((manifests / "details.json").read_text())
    attachments_manifest = json.loads((manifests / "attachments.json").read_text())
    local_projects = json.loads((manifests / "projects.json").read_text())
    corpus_bids_by_project: dict[str, set[str]] = defaultdict(set)
    for row in local_projects:
        if row.get("project_no") and row.get("bid_no"):
            corpus_bids_by_project[row["project_no"]].add(row["bid_no"])
    bids_for_project: dict[str, set[str]] = defaultdict(set)
    for row in candidates:
        key = project_key(row)
        for bid in row.get("existing_db_project_bid_nos", "").split(";"):
            if bid:
                bids_for_project[key].add(bid)
        if not key.startswith("UNLINKED:"):
            bids_for_project[key].update(corpus_bids_by_project.get(key, set()))

    candidate_bids = {bid for bids in bids_for_project.values() for bid in bids}
    indexed_by_bid: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in file_index:
        if row.get("bid_no") in candidate_bids:
            indexed_by_bid[row["bid_no"]].append(row)

    hash_cache: dict[str, dict[str, Any]] = {}
    attachment_rows: list[dict[str, Any]] = []
    for key, bids in bids_for_project.items():
        for bid in sorted(bids):
            detail = details.get(bid)
            cost_relevant = construction_bid_relevant(detail)
            for indexed in indexed_by_bid.get(bid, []):
                source_file = indexed["source_file"]
                path = resolve_source_file(project_root, source_file)
                extension = indexed.get("extension", Path(source_file).suffix.lower())
                digest = sha256(path) if path else ""
                scan: dict[str, Any] = {
                    "scan_error": "파일경로 미확인" if path is None else "",
                    "table_count": 0,
                    "quantity_rows_max": 0,
                    "priced_rows_max": 0,
                    "strong_priced_table": 0,
                    "best_sheet": "",
                    "best_sheet_state": "",
                    "best_header_row": "",
                    "best_priced_ratio": 0.0,
                    "table_signals": [],
                }
                if path and extension in SPREADSHEET_EXTENSIONS:
                    reused = digest in hash_cache
                    if digest not in hash_cache:
                        hash_cache[digest] = scan_spreadsheet(path)
                    scan = hash_cache[digest]
                else:
                    reused = False
                attachment_rows.append(
                    {
                        "project_key": key,
                        "bid_no": bid,
                        "bid_contract_type": clean(detail.get("contract_type")) if detail else "",
                        "bid_title": clean(detail.get("title")) if detail else "",
                        "construction_bid_cost_relevant": int(cost_relevant),
                        "source_file": source_file,
                        "extension": extension,
                        "bytes": indexed.get("bytes", 0),
                        "sha256": digest,
                        "duplicate_scan_reused": int(reused),
                        "boq_filename_candidate": int(bool(BOQ_FILE_RE.search(source_file))),
                        "blank_or_unpriced_filename": int(bool(BLANK_FILE_RE.search(source_file))),
                        **{field: value for field, value in scan.items() if field != "table_signals"},
                        "table_signals_json": json.dumps(scan.get("table_signals", []), ensure_ascii=False),
                    }
                )

    candidates_by_project: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in candidates:
        candidates_by_project[project_key(row)].append(row)
    files_by_project: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in attachment_rows:
        files_by_project[row["project_key"]].append(row)

    project_rows = [
        grade_project(
            key,
            rows,
            files_by_project.get(key, []),
            bids_for_project.get(key, set()),
        )
        for key, rows in candidates_by_project.items()
    ]
    project_rows.sort(key=lambda row: (row["country_ko"], row["grade"], row["project_key"]))
    project_by_key = {row["project_key"]: row for row in project_rows}

    candidate_rows: list[dict[str, Any]] = []
    for row in candidates:
        project = project_by_key[project_key(row)]
        effective_project_no = project["project_no"]
        record_status = record_cost_status(row, project)
        candidate_rows.append(
            {
                "candidate_no": len(candidate_rows) + 1,
                "project_best_grade": project["grade"],
                "project_technical_status": project["technical_status"],
                "record_cost_status": record_status,
                "grade_detail": project["grade_detail"],
                "scope_status": project["scope_status"],
                "price_stage": project["price_stage"],
                "verification_level": project["verification_level"],
                "review_status": project["review_status"],
                "direct_project_estimate_ready": project["direct_project_estimate_ready"],
                "needs_local_rates": project["needs_local_rates"],
                "source_dataset": row["source_dataset"],
                "source_record_id": row["source_record_id"],
                "project_no_linked": row["project_no_linked"],
                "project_no_effective": effective_project_no,
                "project_link_method": (
                    "verified_title_country_alias"
                    if row["source_record_id"] in SOURCE_RECORD_PROJECT_ALIASES
                    else row["project_link_method"]
                ),
                "project_link_score": row["project_link_score"],
                "country_ko": row["country_ko"] or "미상",
                "record_title": row["record_title"],
                "record_scope": row["record_scope"],
                "procurement_category": row["procurement_category"],
                "amount_currency": row["amount_currency"],
                "amount_value": number(row["amount_value"]),
                "amount_role": row["amount_role"],
                "gross_floor_area_m2": number(row["gross_floor_area_m2"]),
                "raw_ceiling_per_gfa_krw": number(row["raw_ceiling_per_gfa_krw"]),
                "linked_bid_count": project["linked_bid_count"],
                "linked_bid_nos": project["linked_bid_nos"],
                "indexed_attachment_count": project["indexed_attachment_count"],
                "spreadsheet_count": project["spreadsheet_count"],
                "boq_named_file_count": project["boq_named_file_count"],
                "quantity_table_file_count": project["quantity_table_file_count"],
                "priced_table_file_count": project["priced_table_file_count"],
                "verified_cost_amount": project["verified_cost_amount"],
                "verified_cost_currency": project["verified_cost_currency"],
                "verified_cost_per_gfa_m2": project["verified_cost_per_gfa_m2"],
                "verified_item_note": project["verified_item_note"],
                "evidence_file": project["evidence_file"],
                "evidence_location": project["evidence_location"],
                "blocking_issue": project["blocking_issue"],
                "next_action": project["next_action"],
                "official_record_url": row["official_record_url"],
                "claim_boundary": project["claim_boundary"],
            }
        )

    countries: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for project in project_rows:
        countries[project["country_ko"] or "미상"].append(project)
    country_rows: list[dict[str, Any]] = []
    for country, projects in countries.items():
        counts = Counter(project["grade"] for project in projects)
        status_counts = Counter(project["technical_status"][0] for project in projects)
        best = "A" if counts["A"] else "B" if counts["B"] else "C"
        if counts["A"] >= 3:
            sufficiency = "초기 국가모형 가능(독립성·시설유형 추가검증 필요)"
        elif counts["A"]:
            sufficiency = "가격사례 부족: 사업별 기준점만 가능"
        elif counts["B"]:
            sufficiency = "수량사례만 있음: 현지단가 필요"
        else:
            sufficiency = "직접 산정근거 없음: 문서추적·원단위 스크리닝만"
        country_rows.append(
            {
                "country_ko": country,
                "country_best_project_grade": best,
                "country_sample_sufficiency": sufficiency,
                "candidate_project_count": len(projects),
                "grade_a_projects": counts["A"],
                "grade_b_projects": counts["B"],
                "grade_c_projects": counts["C"],
                "screening_c_projects": status_counts["C"],
                "unresolved_u_projects": status_counts["U"],
                "direct_estimate_project_count": counts["A"],
                "quantity_only_project_count": counts["B"],
                "linked_bid_count": sum(project["linked_bid_count"] for project in projects),
                "indexed_attachment_count": sum(project["indexed_attachment_count"] for project in projects),
                "note": "최고등급은 해당 국가의 최소 1개 사업 근거이며 국가 대표단가를 뜻하지 않음",
            }
        )
    country_rows.sort(key=lambda row: (row["country_best_project_grade"], -row["grade_a_projects"], row["country_ko"]))

    output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(output_dir / "KOICA_건축후보_127건_등급.csv", candidate_rows)
    write_csv(output_dir / "KOICA_건축후보_프로젝트등급.csv", project_rows)
    write_csv(output_dir / "KOICA_건축후보_국가별준비도.csv", country_rows)
    write_csv(output_dir / "KOICA_건축후보_첨부검사.csv", attachment_rows)

    grade_counts = Counter(row["project_best_grade"] for row in candidate_rows)
    record_status_counts = Counter(row["record_cost_status"][0] for row in candidate_rows)
    project_grade_counts = Counter(row["grade"] for row in project_rows)
    current_candidate_bids = {bid for bid in candidate_bids if bid in details}
    historical_reference_bids = candidate_bids - current_candidate_bids
    projects_with_current_bids = {
        key
        for key, bids in bids_for_project.items()
        if any(bid in current_candidate_bids for bid in bids)
    }
    projects_with_indexed_files = {
        row["project_key"] for row in attachment_rows
    }
    connected_candidate_rows = [
        row for row in candidates if project_key(row) in projects_with_current_bids
    ]
    attachment_candidate_rows = [
        row for row in candidates if project_key(row) in projects_with_indexed_files
    ]
    top_level_attachment_count = sum(
        len(attachments_manifest.get(bid, {}).get("files", []))
        for bid in current_candidate_bids
    )
    summary = {
        "schema_version": "1.0.0",
        "candidate_records": len(candidate_rows),
        "unique_candidate_projects": len(project_rows),
        "countries_including_unknown": len(country_rows),
        "candidate_record_grade_counts": dict(sorted(grade_counts.items())),
        "candidate_record_cost_status_counts": dict(sorted(record_status_counts.items())),
        "project_grade_counts": dict(sorted(project_grade_counts.items())),
        "candidate_bids": len(candidate_bids),
        "current_candidate_bids": len(current_candidate_bids),
        "historical_reference_bids_without_current_corpus": len(historical_reference_bids),
        "current_bids_with_indexed_files": len({row["bid_no"] for row in attachment_rows}),
        "connected_candidate_records": len(connected_candidate_rows),
        "connected_candidate_projects": len(projects_with_current_bids),
        "connected_countries": len({row["country_ko"] or "미상" for row in connected_candidate_rows}),
        "attachment_candidate_records": len(attachment_candidate_rows),
        "attachment_candidate_projects": len(projects_with_indexed_files),
        "attachment_countries": len({row["country_ko"] or "미상" for row in attachment_candidate_rows}),
        "top_level_candidate_attachments": top_level_attachment_count,
        "indexed_candidate_attachments": len(attachment_rows),
        "unique_scanned_spreadsheet_hashes": len(hash_cache),
        "spreadsheet_scan_errors": sum(bool(row["scan_error"]) for row in attachment_rows if row["extension"] in SPREADSHEET_EXTENSIONS),
        "grade_definition": {
            "A": "품목별 수량과 양수 단가/금액을 결합해 사업별 직접 추정 가능",
            "B": "품목별 수량은 있으나 usable 단가가 없어 현지단가 입력 필요",
            "C": "면적·상한·범위·문서추적 수준이며 직접 수량×단가 산정 불가",
        },
        "safeguard_status_definition": {
            "U": "건축 관련 후보이나 직접 산정요건 미충족 또는 첨부·범위 추가확인 필요",
            "X": "설계·감리·물품·전체사업비 등 건축공사비 표본에서 제외",
        },
        "country_boundary": "국가 최고등급은 한 사업의 준비도이며 국가 대표단가 또는 충분한 표본을 뜻하지 않는다.",
        "source_candidate_csv": str(candidates_path.relative_to(project_root)),
        "source_file_index": "data/manifests/file_index.json",
    }
    (output_dir / "KOICA_건축후보_등급요약.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2)
    )
    (output_dir / "KOICA_건축후보_구조화.json").write_text(
        json.dumps(
            {
                "summary": summary,
                "candidate_records": candidate_rows,
                "projects": project_rows,
                "countries": country_rows,
                "attachments": attachment_rows,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--candidates", type=Path, default=DEFAULT_CANDIDATES)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    project_root = args.project_root.resolve()
    candidates = args.candidates if args.candidates.is_absolute() else project_root / args.candidates
    output_dir = args.output_dir if args.output_dir.is_absolute() else project_root / args.output_dir
    summary = run(project_root, candidates, output_dir)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

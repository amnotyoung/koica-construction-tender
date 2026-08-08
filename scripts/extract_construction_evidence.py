#!/usr/bin/env python3
"""Extract auditable construction evidence from downloaded KOICA attachments."""

from __future__ import annotations

import argparse
import json
import re
import struct
import subprocess
import tempfile
import unicodedata
import zipfile
import zlib
from collections import Counter, deque
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

from openpyxl import load_workbook

COUNTRIES = {
    "가나": "Ghana", "과테말라": "Guatemala", "네팔": "Nepal",
    "도미니카": "Dominican Republic", "라오스": "Laos", "르완다": "Rwanda",
    "모로코": "Morocco", "몽골": "Mongolia", "미얀마": "Myanmar",
    "방글라데시": "Bangladesh", "베트남": "Vietnam", "볼리비아": "Bolivia",
    "세네갈": "Senegal", "스리랑카": "Sri Lanka", "알제리": "Algeria",
    "에콰도르": "Ecuador", "에티오피아": "Ethiopia", "엘살바도르": "El Salvador",
    "요르단": "Jordan", "우간다": "Uganda", "우즈베키스탄": "Uzbekistan",
    "이라크": "Iraq", "인도네시아": "Indonesia", "잠비아": "Zambia",
    "캄보디아": "Cambodia", "케냐": "Kenya", "콜롬비아": "Colombia",
    "키르기즈": "Kyrgyzstan", "타지키스탄": "Tajikistan", "탄자니아": "Tanzania",
    "튀니지": "Tunisia", "파라과이": "Paraguay", "파키스탄": "Pakistan",
    "팔레스타인": "Palestine", "페루": "Peru", "필리핀": "Philippines",
    "온두라스": "Honduras", "코트디부아르": "Côte d’Ivoire",
    "카메룬": "Cameroon", "모잠비크": "Mozambique", "말라위": "Malawi",
    "짐바브웨": "Zimbabwe", "동티모르": "Timor-Leste", "피지": "Fiji",
    "솔로몬": "Solomon Islands", "부탄": "Bhutan", "조지아": "Georgia",
}

COUNTRY_ALIASES = {
    "우즈벡": ("우즈베키스탄", "Uzbekistan"),
}

FACILITIES = {
    "병원": ("병원", "의료원", "보건소", "health center", "hospital", "clinic"),
    "학교·교육시설": ("학교", "교실", "교육센터", "school", "classroom", "training center"),
    "실험실·연구시설": ("실험실", "연구소", "laboratory", "lab "),
    "정부·행정시설": ("청사", "경찰", "행정", "government", "police"),
    "농업시설": ("농장", "농업", "관개", "farm", "agricultur", "irrigation"),
    "상하수·환경시설": ("정수", "하수", "폐기물", "water treatment", "waste"),
    "산업·직업훈련시설": ("직업훈련", "기술교육", "산업", "vocational", "technical center"),
}

CATEGORY_KEYWORDS = {
    "연면적·면적": (
        "연면적", "건축면적", "대지면적", "floor area", "gross area",
        "built-up area", "building area", "square meter", "m²", "㎡",
    ),
    "공사비": (
        "공사비", "건축비", "시공비", "construction cost", "works cost",
        "contract amount", "contract price", "bill of quantities", "boq",
    ),
    "설계비": ("설계비", "design fee", "design cost", "consultancy fee"),
    "감리·CM비": (
        "감리비", "감리용역", "supervision fee", "supervision cost",
        "construction management", "cm fee", "bureau de control",
    ),
    "총사업비·예산": (
        "총사업비", "사업비", "예산", "집행한도", "estimated cost",
        "project budget", "total budget", "ceiling amount",
    ),
    "부가세·예비비": (
        "부가가치세", "vat", "tax", "예비비", "contingency", "reserve",
    ),
    "기간": (
        "공사기간", "용역기간", "construction period", "completion period",
        "duration", "calendar days", "months",
    ),
}

AREA_RE = re.compile(
    r"(?<![\w.])(\d[\d,. ]{0,18})\s*(m(?:2|²)|sq\.?\s*m|square\s*met(?:er|re)s?|㎡)",
    re.I,
)
CURRENCY_RE = re.compile(
    r"(?:(USD|US\$|(?<![A-Za-z0-9_])\$|EUR|€|KRW|원|KES|UGX|RWF|XOF|XAF|PKR|PEN|"
    r"IDR|VND|MNT|IQD|DZD|ETB|GHS|TZS|ZMW|MAD|JOD|NPR|BDT)\s*)"
    r"(\d[\d,.\s]{1,22})(?:\s*(million|billion|백만|천))?"
    r"|(\d[\d,.\s]{1,22})\s*"
    r"(USD|US\$|\$|EUR|€|KRW|원|KES|UGX|RWF|XOF|XAF|PKR|PEN|"
    r"IDR|VND|MNT|IQD|DZD|ETB|GHS|TZS|ZMW|MAD|JOD|NPR|BDT)",
    re.I,
)
PERCENT_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*%")

ROOT_ARCHIVE_SUFFIXES = {".zip", ".hwpx"}
NESTED_ARCHIVE_SUFFIXES = {".zip"}
MAX_ARCHIVE_MEMBERS = 5_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 2_000_000_000
MAX_ARCHIVE_TREE_UNCOMPRESSED_BYTES = 2_000_000_000
MAX_NESTED_ARCHIVE_DEPTH = 8
MAX_NESTED_ARCHIVES_PER_TREE = 1_000
UNPACK_MARKER_VERSION = 1
COPY_BUFFER_BYTES = 1024 * 1024


def clean_text(value: Any) -> str:
    text = unicodedata.normalize("NFC", str(value or ""))
    text = text.replace("\x00", " ")
    return re.sub(r"\s+", " ", text).strip()


def safe_member_path(name: str) -> Path | None:
    path = Path(unicodedata.normalize("NFC", name.replace("\\", "/")))
    if not path.parts or path.is_absolute() or ".." in path.parts:
        return None
    return path


def archive_signature(archive: Path) -> dict[str, int]:
    stat = archive.stat()
    return {
        "archive_size": stat.st_size,
        "archive_mtime_ns": stat.st_mtime_ns,
    }


def marker_status(marker: Path, archive: Path, allow_legacy: bool) -> str | None:
    if not marker.is_file():
        return None
    try:
        value = json.loads(marker.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        # Previous versions created empty markers. Trust them only for top-level
        # archives; an archive member named ".complete" must not spoof a nested
        # extraction marker.
        if allow_legacy:
            try:
                return "legacy" if marker.stat().st_size == 0 else None
            except OSError:
                return None
        return None
    if not isinstance(value, dict) or value.get("version") != UNPACK_MARKER_VERSION:
        return None
    signature = archive_signature(archive)
    if all(value.get(key) == expected for key, expected in signature.items()):
        return "versioned"
    return None


def write_unpack_marker(marker: Path, archive: Path) -> None:
    value = {"version": UNPACK_MARKER_VERSION, **archive_signature(archive)}
    marker.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        prefix=".complete-",
        suffix=".tmp",
        dir=marker.parent,
        delete=False,
    ) as handle:
        json.dump(value, handle, ensure_ascii=False, sort_keys=True)
        handle.write("\n")
        temp_marker = Path(handle.name)
    temp_marker.replace(marker)


def path_within(path: Path, root: Path) -> bool:
    try:
        path.resolve(strict=False).relative_to(root.resolve(strict=False))
        return True
    except ValueError:
        return False


def validated_members(zf: zipfile.ZipFile) -> tuple[list[zipfile.ZipInfo], int]:
    members = zf.infolist()
    if len(members) > MAX_ARCHIVE_MEMBERS:
        raise ValueError(f"archive has more than {MAX_ARCHIVE_MEMBERS:,} members")
    total = sum(item.file_size for item in members)
    if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise ValueError(
            "archive expands beyond "
            f"{MAX_ARCHIVE_UNCOMPRESSED_BYTES:,} bytes"
        )
    return members, total


def extract_archive(
    archive: Path,
    target: Path,
    members: Iterable[zipfile.ZipInfo],
    zf: zipfile.ZipFile,
) -> None:
    if target.is_symlink():
        raise ValueError(f"archive target is a symlink: {target}")
    target.mkdir(parents=True, exist_ok=True)
    target_root = target.resolve()
    extracted_bytes = 0
    for item in members:
        member = safe_member_path(item.filename)
        if member is None or item.is_dir():
            continue
        destination = target / member
        if not path_within(destination, target_root):
            raise ValueError(f"archive member escapes target: {item.filename}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        # Resolve again after mkdir so a pre-existing symlink in a partially
        # extracted directory cannot redirect writes outside the target.
        if not path_within(destination, target_root) or destination.is_symlink():
            raise ValueError(f"unsafe archive destination: {item.filename}")
        item_bytes = 0
        with zf.open(item) as src, destination.open("wb") as dst:
            while True:
                chunk = src.read(COPY_BUFFER_BYTES)
                if not chunk:
                    break
                item_bytes += len(chunk)
                extracted_bytes += len(chunk)
                if (
                    item_bytes > item.file_size
                    or extracted_bytes > MAX_ARCHIVE_UNCOMPRESSED_BYTES
                ):
                    raise ValueError("archive expanded beyond declared limits")
                dst.write(chunk)


def is_versioned_unpack_dir(path: Path) -> bool:
    marker = path / ".complete"
    if not marker.is_file():
        return False
    try:
        value = json.loads(marker.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return False
    return (
        isinstance(value, dict)
        and value.get("version") == UNPACK_MARKER_VERSION
    )


def nested_archives(target: Path) -> list[Path]:
    """Return ZIPs owned by target, excluding already unpacked child trees."""
    result: list[Path] = []
    if not target.is_dir():
        return result
    for archive in sorted(target.rglob("*"), key=lambda path: str(path).casefold()):
        if (
            not archive.is_file()
            or archive.is_symlink()
            or archive.suffix.lower() not in NESTED_ARCHIVE_SUFFIXES
        ):
            continue
        ancestor = archive.parent
        belongs_to_child_tree = False
        while ancestor != target:
            if is_versioned_unpack_dir(ancestor):
                belongs_to_child_tree = True
                break
            if ancestor.parent == ancestor:
                break
            ancestor = ancestor.parent
        if not belongs_to_child_tree:
            result.append(archive)
    return result


def unpack_archives(raw_dir: Path, unpacked_dir: Path) -> dict[str, int]:
    stats = Counter()
    queue: deque[tuple[Path, Path, int, int]] = deque()
    queued: set[tuple[str, str]] = set()
    tree_bytes: dict[int, int] = {}
    tree_nested_counts: Counter[int] = Counter()

    root_archives = sorted(raw_dir.rglob("*"), key=lambda path: str(path).casefold())
    for archive in root_archives:
        if (
            not archive.is_file()
            or archive.is_symlink()
            or archive.suffix.lower() not in ROOT_ARCHIVE_SUFFIXES
        ):
            continue
        rel = archive.relative_to(raw_dir)
        target = unpacked_dir / rel.with_suffix("")
        tree_id = len(tree_bytes)
        tree_bytes[tree_id] = 0
        key = (str(archive.resolve()), str(target.resolve(strict=False)))
        queued.add(key)
        queue.append((archive, target, 0, tree_id))

    while queue:
        archive, target, depth, tree_id = queue.popleft()
        if not path_within(target, unpacked_dir):
            stats["failed"] += 1
            continue
        marker = target / ".complete"
        status = marker_status(marker, archive, allow_legacy=depth == 0)
        if status:
            stats["reused"] += 1
            if depth:
                stats["nested_reused"] += 1
            if status == "legacy":
                write_unpack_marker(marker, archive)
                stats["legacy_markers_upgraded"] += 1
        else:
            try:
                with zipfile.ZipFile(archive) as zf:
                    members, total = validated_members(zf)
                    if (
                        tree_bytes[tree_id] + total
                        > MAX_ARCHIVE_TREE_UNCOMPRESSED_BYTES
                    ):
                        raise ValueError(
                            "archive tree expands beyond "
                            f"{MAX_ARCHIVE_TREE_UNCOMPRESSED_BYTES:,} bytes"
                        )
                    # Reserve the budget before writing. A failed partial
                    # extraction must not regain its consumed safety allowance.
                    tree_bytes[tree_id] += total
                    extract_archive(archive, target, members, zf)
                write_unpack_marker(marker, archive)
                stats["unpacked"] += 1
                if depth:
                    stats["nested_unpacked"] += 1
            except Exception:
                stats["failed"] += 1
                if depth:
                    stats["nested_failed"] += 1
                continue

        children = nested_archives(target)
        if not children:
            continue
        if depth >= MAX_NESTED_ARCHIVE_DEPTH:
            stats["depth_limited"] += len(children)
            continue
        for child in children:
            child_target = child.with_suffix("")
            key = (str(child.resolve()), str(child_target.resolve(strict=False)))
            if key in queued:
                continue
            if (
                tree_nested_counts[tree_id]
                >= MAX_NESTED_ARCHIVES_PER_TREE
            ):
                stats["archive_count_limited"] += 1
                continue
            queued.add(key)
            tree_nested_counts[tree_id] += 1
            queue.append((child, child_target, depth + 1, tree_id))
    return dict(stats)


def hwp_chunks(path: Path) -> list[tuple[str, str]]:
    import olefile

    chunks: list[tuple[str, str]] = []
    with olefile.OleFileIO(path) as ole:
        header = ole.openstream("FileHeader").read()
        compressed = bool(struct.unpack("<I", header[36:40])[0] & 1)
        sections = sorted(
            name for name in ole.listdir() if name and name[0] == "BodyText"
        )
        para_no = 0
        for stream_name in sections:
            data = ole.openstream(stream_name).read()
            if compressed:
                data = zlib.decompress(data, -15)
            offset = 0
            while offset + 4 <= len(data):
                record = struct.unpack("<I", data[offset : offset + 4])[0]
                offset += 4
                tag_id = record & 0x3FF
                size = (record >> 20) & 0xFFF
                if size == 0xFFF:
                    size = struct.unpack("<I", data[offset : offset + 4])[0]
                    offset += 4
                payload = data[offset : offset + size]
                offset += size
                if tag_id == 67:
                    para_no += 1
                    text = clean_text(payload.decode("utf-16le", errors="ignore"))
                    if text:
                        chunks.append((f"문단 {para_no}", text))
    return chunks


def pdf_chunks(path: Path) -> list[tuple[str, str]]:
    result = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"],
        check=False,
        capture_output=True,
        timeout=180,
    )
    text = result.stdout.decode("utf-8", errors="replace")
    return [
        (f"페이지 {index}", clean_text(page))
        for index, page in enumerate(text.split("\f"), 1)
        if clean_text(page)
    ]


def docx_chunks(path: Path) -> list[tuple[str, str]]:
    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    body = root.find(f"{{{word_namespace}}}body")
    if body is None:
        return []
    chunks: list[tuple[str, str]] = []
    paragraph_no = 0
    table_no = 0
    for child in body:
        if child.tag == f"{{{word_namespace}}}p":
            value = clean_text(
                " ".join(
                    node.text or ""
                    for node in child.iter(f"{{{word_namespace}}}t")
                )
            )
            if value:
                paragraph_no += 1
                chunks.append((f"문단 {paragraph_no}", value))
        elif child.tag == f"{{{word_namespace}}}tbl":
            table_no += 1
            rows = child.findall(f"{{{word_namespace}}}tr")
            for row_no, row in enumerate(rows, 1):
                cells = []
                for cell in row.findall(f"{{{word_namespace}}}tc"):
                    cells.append(
                        clean_text(
                            " ".join(
                                node.text or ""
                                for node in cell.iter(f"{{{word_namespace}}}t")
                            )
                        )
                    )
                value = clean_text(" | ".join(cells))
                if value:
                    chunks.append((f"표 {table_no} 행 {row_no}", value))
    return chunks


def xlsx_chunks(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    chunks: list[tuple[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(), 1):
                values = [
                    f"{cell.coordinate}={clean_text(cell.value)}"
                    for cell in row
                    if cell.value not in (None, "")
                ]
                if values:
                    chunks.append((f"{sheet.title}!{row_number}", " | ".join(values)))
    finally:
        workbook.close()
    return chunks


def xls_chunks(path: Path) -> list[tuple[str, str]]:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    chunks: list[tuple[str, str]] = []
    try:
        for sheet in workbook.sheets():
            for row_number in range(sheet.nrows):
                values = []
                for column_number in range(sheet.ncols):
                    value = clean_text(sheet.cell_value(row_number, column_number))
                    if value:
                        values.append(f"R{row_number + 1}C{column_number + 1}={value}")
                if values:
                    chunks.append((f"{sheet.name}!{row_number + 1}", " | ".join(values)))
    finally:
        workbook.release_resources()
    return chunks


def pptx_chunks(path: Path) -> list[tuple[str, str]]:
    drawing_namespace = (
        "http://schemas.openxmlformats.org/drawingml/2006/main"
    )
    chunks: list[tuple[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        slide_names = sorted(
            (
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=lambda name: int(re.search(r"\d+", Path(name).stem).group()),
        )
        for slide_number, name in enumerate(slide_names, 1):
            root = ElementTree.fromstring(archive.read(name))
            values = [
                clean_text(node.text)
                for node in root.iter(f"{{{drawing_namespace}}}t")
                if clean_text(node.text)
            ]
            if values:
                chunks.append((f"슬라이드 {slide_number}", " | ".join(values)))
    return chunks


def legacy_doc_chunks(path: Path) -> list[tuple[str, str]]:
    with tempfile.TemporaryDirectory(prefix="koica-doc-") as temp:
        output = Path(temp) / "converted.txt"
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-output", str(output), str(path)],
            check=False,
            capture_output=True,
            timeout=180,
        )
        if result.returncode or not output.exists():
            raise RuntimeError(result.stderr.decode("utf-8", errors="replace"))
        text = output.read_text(encoding="utf-8", errors="replace")
    return [
        (f"문단 {index}", clean_text(paragraph))
        for index, paragraph in enumerate(re.split(r"\n\s*\n", text), 1)
        if clean_text(paragraph)
    ]


def xml_chunks(path: Path) -> list[tuple[str, str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    text = clean_text(re.sub(r"<[^>]+>", " ", raw))
    return [("본문", text)] if text else []


def extract_chunks(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return pdf_chunks(path)
    if suffix == ".docx":
        return docx_chunks(path)
    if suffix in {".xlsx", ".xlsm"}:
        return xlsx_chunks(path)
    if suffix == ".xls":
        return xls_chunks(path)
    if suffix == ".pptx":
        return pptx_chunks(path)
    if suffix == ".doc":
        return legacy_doc_chunks(path)
    if suffix == ".hwp":
        return hwp_chunks(path)
    if suffix in {".txt", ".csv", ".xml", ".html", ".htm"}:
        return xml_chunks(path)
    return []


def is_source_document(relative_path: Path) -> bool:
    """Keep meaningful dot-directories while excluding generated metadata."""
    if not relative_path.parts:
        return False
    if "__MACOSX" in relative_path.parts:
        return False
    name = relative_path.name
    return not (
        name in {".complete", ".DS_Store", "Thumbs.db"}
        or name.startswith(".complete-")
        or name.startswith("._")
        or name.startswith("~$")
    )


def find_country(text: str) -> tuple[str, str]:
    lowered = text.lower()
    matches: list[tuple[int, str, str]] = []
    for korean, english in COUNTRIES.items():
        if korean in text or english.lower() in lowered:
            matched_length = max(
                len(marker)
                for marker in (korean, english)
                if marker in text or marker.lower() in lowered
            )
            matches.append((matched_length, korean, english))
    for alias, (korean, english) in COUNTRY_ALIASES.items():
        if alias in text:
            matches.append((len(alias), korean, english))
    if not matches:
        return "", ""
    _, korean, english = max(matches, key=lambda match: match[0])
    return korean, english


def classify_facility(text: str) -> str:
    lowered = text.lower()
    for label, keywords in FACILITIES.items():
        if any(keyword in lowered for keyword in keywords):
            return label
    return "기타·미분류"


def classify_work(text: str) -> str:
    lowered = text.lower()
    if any(k in lowered for k in ("리모델링", "개보수", "보수", "renovation", "rehabilitation")):
        return "리모델링·개보수"
    if any(k in lowered for k in ("증축", "extension", "expansion")):
        return "증축"
    if any(k in lowered for k in ("신축", "건립", "new construction", "construction of")):
        return "신축"
    return "미분류"


def numeric_tokens(text: str) -> dict[str, list[str]]:
    return {
        "areas": [clean_text(" ".join(m.groups(default=""))) for m in AREA_RE.finditer(text)],
        "currencies": [
            clean_text(" ".join(m.groups(default=""))) for m in CURRENCY_RE.finditer(text)
        ],
        "percentages": [m.group(0) for m in PERCENT_RE.finditer(text)],
    }


def evidence_from_chunks(
    bid_no: str, source_path: Path, chunks: Iterable[tuple[str, str]], root: Path
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    relative = str(source_path.relative_to(root))
    for locator, text in chunks:
        lowered = text.lower()
        tokens = numeric_tokens(text)
        if not any(tokens.values()):
            continue
        for category, keywords in CATEGORY_KEYWORDS.items():
            matched = [keyword for keyword in keywords if keyword in lowered]
            if not matched:
                continue
            if category == "연면적·면적" and not tokens["areas"]:
                continue
            if category in {"공사비", "설계비", "감리·CM비", "총사업비·예산"} and not tokens["currencies"]:
                continue
            records.append(
                {
                    "bid_no": bid_no,
                    "category": category,
                    "matched_keywords": ", ".join(matched[:5]),
                    "area_mentions": " | ".join(tokens["areas"][:8]),
                    "currency_mentions": " | ".join(tokens["currencies"][:8]),
                    "percentage_mentions": " | ".join(tokens["percentages"][:8]),
                    "source_file": relative,
                    "source_locator": locator,
                    "evidence_text": text[:1200],
                    "review_status": "자동추출·검토필요",
                }
            )
    return records


def project_rows(details: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for bid_no, detail in sorted(details.items(), reverse=True):
        fields = detail.get("fields", {})
        combined = " ".join(
            (
                fields.get("사업명", ""),
                fields.get("입찰명(국문)", ""),
                fields.get("입찰명(영문)", ""),
            )
        )
        country_ko, country_en = find_country(combined)
        rows.append(
            {
                "bid_no": bid_no,
                "project_no": fields.get("사업번호", ""),
                "country_ko": country_ko,
                "country_en": country_en,
                "region": "",
                "project_name": fields.get("사업명", ""),
                "bid_title_ko": fields.get("입찰명(국문)", ""),
                "bid_title_en": fields.get("입찰명(영문)", ""),
                "facility_type": classify_facility(combined),
                "work_type": classify_work(combined),
                "contract_type": fields.get("계약구분", ""),
                "contract_method": fields.get("계약방법", ""),
                "selection_method": fields.get("낙찰자 선정방식", ""),
                "ceiling_usd_raw": fields.get("집행한도금액(달러)", ""),
                "ceiling_krw_raw": fields.get("집행한도금액(원)", ""),
                "notice_date": fields.get("공고일자", ""),
                "detail_url": detail.get("detail_url", ""),
            }
        )
    return rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", type=Path, default=Path("data"))
    parser.add_argument("--max-files", type=int)
    args = parser.parse_args()
    data = args.data.resolve()
    raw = data / "raw"
    unpacked = data / "unpacked"
    manifests = data / "manifests"
    unpack_stats = unpack_archives(raw, unpacked)
    details = json.loads((manifests / "details.json").read_text(encoding="utf-8"))

    sources: list[tuple[str, Path, Path]] = []
    for root in (raw, unpacked):
        for path in root.rglob("*"):
            if not path.is_file() or path.name == ".complete":
                continue
            rel = path.relative_to(root)
            # Office lock files and OS/archive metadata are not source documents.
            if not is_source_document(rel):
                continue
            sources.append((rel.parts[0], path, root))
    if args.max_files:
        sources = sources[: args.max_files]

    evidence: list[dict[str, Any]] = []
    file_index: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    for index, (bid_no, path, root) in enumerate(sources, 1):
        try:
            chunks = extract_chunks(path)
            found = evidence_from_chunks(bid_no, path, chunks, root)
            evidence.extend(found)
            file_index.append(
                {
                    "bid_no": bid_no,
                    "source_file": str(path.relative_to(root)),
                    "extension": path.suffix.lower(),
                    "bytes": path.stat().st_size,
                    "text_chunks": len(chunks),
                    "evidence_count": len(found),
                }
            )
        except Exception as exc:
            errors.append(
                {
                    "bid_no": bid_no,
                    "source_file": str(path.relative_to(root)),
                    "error": str(exc),
                }
            )
        if index % 100 == 0:
            print(f"files {index}/{len(sources)} evidence={len(evidence)}", flush=True)

    outputs = {
        "projects.json": project_rows(details),
        "construction_evidence.json": evidence,
        "file_index.json": file_index,
        "extraction_errors.json": errors,
        "extraction_summary.json": {
            "unpack": unpack_stats,
            "source_files": len(sources),
            "indexed_files": len(file_index),
            "evidence_records": len(evidence),
            "errors": len(errors),
        },
    }
    for name, value in outputs.items():
        (manifests / name).write_text(
            json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(json.dumps(outputs["extraction_summary.json"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

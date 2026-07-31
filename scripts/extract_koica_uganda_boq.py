#!/usr/bin/env python3
"""Extract an auditable KOICA Uganda BOQ into normalized row-level records.

The source is an unpriced KOICA bidding workbook. The extractor preserves the
original sheet/row lineage, cached quantity values, and Excel formulas while
adding conservative facility, trade, unit, and quality-control fields.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_SOURCE = Path(
    "data/unpacked/L2023-00003-1/"
    "Bid Document(ESMV II)_separately attached/"
    "Bid Document(ESMV II)_separately attached/"
    "Part 2 Requirements/"
    "Section 4. Bill of Quantities_unpriced.xlsx"
)

SHEET_SCOPE = {
    "PRELIMINARIES": ("공통가설·일반조건", 1, 1, "세부시트 1식"),
    "EARTHWORKS": ("부지 토공", 1, 1, "General Summary 수량 1"),
    "2-STANCE PIT LATRINE": (
        "2칸 재래식 화장실",
        1,
        1,
        "General Summary 수량 1",
    ),
    "6-STANCE WATER BORNE TOILET": (
        "6칸 수세식 화장실",
        1,
        1,
        "General Summary 수량 1",
    ),
    "MILLING FACILITY": ("옥수수 가공시설", 1, 1, "General Summary 수량 1"),
    "ADMIN BLOCK": ("관리동", 1, 1, "General Summary 수량 1"),
    "MARKET SHED": (
        "여성친화형 시장 쉼터",
        5,
        6,
        "세부시트는 1동, General Summary는 5동이나 재공고 전기공종 "
        "제목은 6동을 명시하여 범위 확인 필요",
    ),
    "RAINWATER HARVESTING": (
        "빗물저장시설 10,000L",
        2,
        2,
        "세부시트는 1기, General Summary 수량 2",
    ),
    "FIRE SUPRESSION SYSTEM": ("소화설비", 1, 1, "General Summary 수량 1"),
    "EXTERNAL WORKS": ("외부공사", 1, 1, "General Summary 수량 1"),
    "SUPPLY OF 33KV LINE": (
        "33kV 전력인입",
        1,
        1,
        "General Summary 수량 1",
    ),
}

UNIT_MAP = {
    "SM": "m²",
    "SQM": "m²",
    "M2": "m²",
    "M²": "m²",
    "CM": "m³",
    "CUM": "m³",
    "M3": "m³",
    "M³": "m³",
    "LM": "m",
    "L.M": "m",
    "L/M": "m",
    "METRE": "m",
    "METRES": "m",
    "METER": "m",
    "METERS": "m",
    "KG": "kg",
    "KGS": "kg",
    "TON": "t",
    "TONNE": "t",
    "NO": "no.",
    "NO.": "no.",
    "NR": "no.",
    "EACH": "no.",
    "EA": "no.",
    "PCS": "no.",
    "PIECE": "no.",
    "PIECES": "no.",
    "PC": "no.",
    "SET": "set",
    "SETS": "set",
    "PKT": "pkt",
    "PKTS": "pkt",
    "ITEM": "item",
    "SUM": "item",
    "LS": "item",
    "LOT": "item",
    "L": "L",
    "LTR": "L",
    "LITRE": "L",
    "LITRES": "L",
}

TOTAL_RE = re.compile(
    r"\b(?:sub\s*total|subtotal|grand\s+total|total|carried\s+to|"
    r"brought\s+forward|collection)\b",
    re.IGNORECASE,
)
ELEMENT_RE = re.compile(r"\belement\s*(?:no\.?)?\s*\d+", re.IGNORECASE)


@dataclass
class Header:
    row: int
    item_col: int | None
    description_col: int
    unit_col: int | None
    quantity_col: int | None
    rate_col: int | None
    amount_col: int | None


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def infer_bid_no(path: Path) -> str:
    match = re.search(r"L\d{4}-\d{5}-\d", str(path))
    return match.group(0) if match else ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def number_or_none(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    return None


def normalize_unit(value: Any) -> str:
    raw = clean_text(value).upper().replace(" ", "")
    return UNIT_MAP.get(raw, clean_text(value))


def facility_scope(sheet_name: str) -> tuple[str, int, int, str] | None:
    normalized = clean_text(sheet_name).upper()
    for key, scope in SHEET_SCOPE.items():
        if key in normalized:
            return scope
    return None


def find_header(sheet) -> Header | None:
    if "RAINWATER HARVESTING" in clean_text(sheet.title).upper():
        return Header(
            row=1,
            item_col=1,
            description_col=2,
            unit_col=3,
            quantity_col=4,
            rate_col=5,
            amount_col=6,
        )
    for row_no in range(1, min(sheet.max_row, 12) + 1):
        cells = {
            col: clean_text(sheet.cell(row_no, col).value).lower()
            for col in range(1, min(sheet.max_column, 20) + 1)
        }
        description_col = next(
            (
                col
                for col, value in cells.items()
                if value in {"description", "item description", "bill summary"}
            ),
            None,
        )
        if description_col is None:
            continue
        item_col = next(
            (
                col
                for col, value in cells.items()
                if value in {"s/n", "item", "item no:", "item no"}
            ),
            None,
        )
        unit_col = next(
            (col for col, value in cells.items() if value in {"unit", "uom"}),
            None,
        )
        quantity_col = next(
            (
                col
                for col, value in cells.items()
                if value in {"qty", "quantity", "quantity "}
            ),
            None,
        )
        rate_col = next(
            (col for col, value in cells.items() if "rate" in value),
            None,
        )
        amount_col = next(
            (
                col
                for col, value in cells.items()
                if "amount" in value or value == "total"
            ),
            None,
        )
        return Header(
            row=row_no,
            item_col=item_col,
            description_col=description_col,
            unit_col=unit_col,
            quantity_col=quantity_col,
            rate_col=rate_col,
            amount_col=amount_col,
        )
    return None


def classify_trade(
    sheet_name: str,
    description: str,
    element_context: str,
    heading_context: str,
) -> str:
    haystack = " ".join(
        (sheet_name, element_context, heading_context, description)
    ).lower()
    sheet = sheet_name.lower()
    if "preliminar" in sheet:
        return "가설·일반조건"
    if "33kv" in sheet:
        return "전기·전력인입"
    if "fire" in sheet:
        return "소방"
    if "rainwater" in sheet:
        return "급배수·빗물이용"
    if "external works" in sheet:
        return "외부공사"

    rules = (
        ("소방", r"fire|extinguisher|suppression"),
        (
            "전기·전력인입",
            r"electrical|wiring|light(?:ing)?|socket|cable|breaker|"
            r"switch|consumer unit|photocell|pole|transformer|33\s*kv",
        ),
        (
            "급배수·위생",
            r"sanitary|plumb|water closet|\bwc\b|wash basin|pipe|"
            r"drain|sewer|septic|manhole|tap|shower|toilet",
        ),
        (
            "지붕",
            r"roof|truss|purlin|rafter|fascia|facia|gutter|downpipe|"
            r"rainwater disposal",
        ),
        ("창호·철물", r"door|window|glaz|ironmongery|lock|hinge"),
        (
            "마감",
            r"finish|plaster|render|paint|screed|tile|ceiling|roughcast|"
            r"terrazzo|vinyl",
        ),
        (
            "조적",
            r"brick|block|walling|mortar|masonry|damp proof course",
        ),
        (
            "콘크리트·철근·거푸집",
            r"concrete|reinforcement|reinforced|steel bar|mesh|formwork|"
            r"ringbeam|column|beam|slab",
        ),
        (
            "토공·기초",
            r"excavat|earthwork|top\s*soil|hard\s*core|hardcore|marrum|"
            r"anti[- ]termite|foundation|substructure|fill and ram",
        ),
        (
            "외부공사",
            r"fenc|gate|external works|road|parking|landscap|paving|walkway",
        ),
        (
            "설비·장비",
            r"mill(?:ing)?|machine|equipment|mechanical|ventilat|air condition",
        ),
        (
            "가설·일반조건",
            r"preliminar|insurance|site office|security|progress photograph|"
            r"setting[- ]out|scaffold|signboard|health and safety",
        ),
    )
    for label, pattern in rules:
        if re.search(pattern, haystack):
            return label
    return "기타"


def classify_row(
    description: str,
    item_code: str,
    unit_raw: str,
    quantity: float | None,
    quantity_formula: str,
    is_preliminaries: bool,
) -> str:
    if TOTAL_RE.search(description):
        return "total"
    if unit_raw and (quantity is not None or quantity_formula):
        return "item"
    if item_code and description and is_preliminaries:
        return "item_unquantified"
    if item_code and description and unit_raw:
        return "item_unquantified"
    return "heading"


def formula_text(cell) -> str:
    value = cell.value
    return value if isinstance(value, str) and value.startswith("=") else ""


def last_meaningful_row(sheet, columns: Iterable[int]) -> int:
    last = 0
    for row_no in range(1, sheet.max_row + 1):
        if any(sheet.cell(row_no, col).value not in (None, "") for col in columns):
            last = row_no
    return last


def extract_boq(source: Path) -> dict[str, Any]:
    values_book = load_workbook(source, data_only=True, read_only=False)
    formulas_book = load_workbook(source, data_only=False, read_only=False)
    records: list[dict[str, Any]] = []
    sheet_stats: list[dict[str, Any]] = []
    skipped_sheets: list[dict[str, str]] = []

    for formula_sheet in formulas_book.worksheets:
        sheet_name = clean_text(formula_sheet.title)
        if "SUMMARY" in sheet_name.upper():
            skipped_sheets.append(
                {"sheet": sheet_name, "reason": "요약 시트(상세 항목과 중복)"}
            )
            continue
        scope = facility_scope(sheet_name)
        if scope is None:
            skipped_sheets.append(
                {"sheet": sheet_name, "reason": "표지·요약·구분 시트"}
            )
            continue
        header = find_header(formula_sheet)
        if header is None:
            skipped_sheets.append(
                {"sheet": sheet_name, "reason": "열 머리글 탐지 실패"}
            )
            continue
        value_sheet = values_book[formula_sheet.title]
        columns = [
            col
            for col in (
                header.item_col,
                header.description_col,
                header.unit_col,
                header.quantity_col,
                header.rate_col,
                header.amount_col,
            )
            if col is not None
        ]
        last_row = last_meaningful_row(formula_sheet, columns)
        facility, multiplier_low, multiplier_high, multiplier_basis = scope
        element_context = ""
        heading_context = ""
        start_index = len(records)
        is_preliminaries = "PRELIMINAR" in sheet_name.upper()

        for row_no in range(header.row + 1, last_row + 1):
            description = clean_text(
                value_sheet.cell(row_no, header.description_col).value
                or formula_sheet.cell(row_no, header.description_col).value
            )
            if not description:
                continue
            item_code = (
                clean_text(
                    value_sheet.cell(row_no, header.item_col).value
                    or formula_sheet.cell(row_no, header.item_col).value
                )
                if header.item_col
                else ""
            )
            unit_raw = (
                clean_text(value_sheet.cell(row_no, header.unit_col).value)
                if header.unit_col
                else ""
            )
            quantity = (
                number_or_none(value_sheet.cell(row_no, header.quantity_col).value)
                if header.quantity_col
                else None
            )
            quantity_formula = (
                formula_text(formula_sheet.cell(row_no, header.quantity_col))
                if header.quantity_col
                else ""
            )
            rate = (
                number_or_none(value_sheet.cell(row_no, header.rate_col).value)
                if header.rate_col
                else None
            )
            rate_formula = (
                formula_text(formula_sheet.cell(row_no, header.rate_col))
                if header.rate_col
                else ""
            )
            amount = (
                number_or_none(value_sheet.cell(row_no, header.amount_col).value)
                if header.amount_col
                else None
            )
            amount_formula = (
                formula_text(formula_sheet.cell(row_no, header.amount_col))
                if header.amount_col
                else ""
            )
            row_type = classify_row(
                description,
                item_code,
                unit_raw,
                quantity,
                quantity_formula,
                is_preliminaries,
            )
            if row_type == "heading":
                if ELEMENT_RE.search(description):
                    element_context = description
                    heading_context = ""
                else:
                    heading_context = description
            normalized_unit = normalize_unit(unit_raw)
            scaled_quantity_low = (
                quantity * multiplier_low if quantity is not None else None
            )
            scaled_quantity_high = (
                quantity * multiplier_high if quantity is not None else None
            )
            qc_flags: list[str] = []
            if row_type == "item" and quantity is None:
                qc_flags.append("missing_quantity")
            if row_type == "item_unquantified":
                qc_flags.append("unquantified_allowance")
            if quantity == 0:
                qc_flags.append("zero_quantity_check_formula_or_design")
            if quantity_formula:
                qc_flags.append("formula_quantity")
            if multiplier_high > 1 and row_type == "item":
                qc_flags.append("scope_multiplier_applied_for_scenario")
            if multiplier_low != multiplier_high and row_type == "item":
                qc_flags.append("scope_multiplier_conflict")
            if (rate is None or rate == 0) and (amount is None or amount == 0):
                qc_flags.append("unpriced_source")

            records.append(
                {
                    "bid_no": infer_bid_no(source),
                    "project_no": "2019-00005",
                    "source_file": str(source),
                    "source_sheet": sheet_name,
                    "source_row": row_no,
                    "facility": facility,
                    "scope_multiplier_scenario_low": multiplier_low,
                    "scope_multiplier_scenario_high": multiplier_high,
                    "scope_multiplier_basis": multiplier_basis,
                    "element_context": element_context,
                    "heading_context": heading_context,
                    "row_type": row_type,
                    "item_code": item_code,
                    "description": description,
                    "trade_group": classify_trade(
                        sheet_name,
                        description,
                        element_context,
                        heading_context,
                    ),
                    "unit_raw": unit_raw,
                    "unit_normalized": normalized_unit,
                    "quantity_base": quantity,
                    "quantity_formula": quantity_formula,
                    "quantity_scaled_low": scaled_quantity_low,
                    "quantity_scaled_high": scaled_quantity_high,
                    "rate_ugx": rate,
                    "rate_formula": rate_formula,
                    "amount_ugx": amount,
                    "amount_formula": amount_formula,
                    "is_priced": bool((rate or 0) > 0 or (amount or 0) > 0),
                    "qc_flag": ";".join(qc_flags),
                }
            )

        sheet_records = records[start_index:]
        item_rows = [
            row
            for row in sheet_records
            if row["row_type"] in {"item", "item_unquantified"}
        ]
        sheet_stats.append(
            {
                "sheet": sheet_name,
                "facility": facility,
                "header_row": header.row,
                "last_meaningful_row": last_row,
                "scope_multiplier_scenario_low": multiplier_low,
                "scope_multiplier_scenario_high": multiplier_high,
                "records": len(sheet_records),
                "item_rows": len(item_rows),
                "quantified_items": sum(
                    row["quantity_base"] is not None for row in item_rows
                ),
                "positive_quantity_items": sum(
                    (row["quantity_base"] or 0) > 0 for row in item_rows
                ),
                "zero_quantity_items": sum(
                    row["quantity_base"] == 0 for row in item_rows
                ),
                "priced_items": sum(row["is_priced"] for row in item_rows),
            }
        )

    values_book.close()
    formulas_book.close()
    item_records = [
        row
        for row in records
        if row["row_type"] in {"item", "item_unquantified"}
    ]
    unit_totals: dict[tuple[str, str], dict[str, Any]] = defaultdict(
        lambda: {
            "base_quantity": 0.0,
            "scaled_low_quantity": 0.0,
            "scaled_high_quantity": 0.0,
            "item_rows": 0,
        }
    )
    for row in item_records:
        if row["quantity_base"] is None or not row["unit_normalized"]:
            continue
        key = (row["trade_group"], row["unit_normalized"])
        unit_totals[key]["base_quantity"] += row["quantity_base"]
        unit_totals[key]["scaled_low_quantity"] += row["quantity_scaled_low"]
        unit_totals[key]["scaled_high_quantity"] += row["quantity_scaled_high"]
        unit_totals[key]["item_rows"] += 1

    trade_counts = Counter(row["trade_group"] for row in item_records)
    summary = {
        "source_file": str(source),
        "source_type": "KOICA local bidding attachment",
        "bid_no": infer_bid_no(source),
        "related_bid_nos": ["L2023-00003-1", "L2023-00009-1"],
        "project_no": "2019-00005",
        "project_name": "우간다 지속가능 농촌개발 시범마을 확산 사업",
        "bid_title": (
            "Construction of Medium Maize Milling Facility and Women Friendly "
            "Market in Bongole Parish, Mpigi District"
        ),
        "all_extracted_rows": len(records),
        "item_rows": len(item_records),
        "quantified_items": sum(
            row["quantity_base"] is not None for row in item_records
        ),
        "positive_quantity_items": sum(
            (row["quantity_base"] or 0) > 0 for row in item_records
        ),
        "zero_quantity_items": sum(
            row["quantity_base"] == 0 for row in item_records
        ),
        "unquantified_items": sum(
            row["quantity_base"] is None for row in item_records
        ),
        "formula_quantity_items": sum(
            bool(row["quantity_formula"]) for row in item_records
        ),
        "priced_items": sum(row["is_priced"] for row in item_records),
        "detail_sheets": len(sheet_stats),
        "facility_components": len({row["facility"] for row in item_records}),
        "trade_item_counts": [
            {"trade_group": trade, "item_rows": count}
            for trade, count in trade_counts.most_common()
        ],
        "unit_totals": [
            {
                "trade_group": trade,
                "unit": unit,
                **{
                    key: round(value, 6) if isinstance(value, float) else value
                    for key, value in totals.items()
                },
            }
            for (trade, unit), totals in sorted(unit_totals.items())
        ],
        "sheet_stats": sheet_stats,
        "skipped_sheets": skipped_sheets,
        "interpretation_limits": [
            "원본은 미가격 BOQ로 단가와 금액이 공란이다.",
            "시장쉼터는 General Summary 5동과 재공고 세부공종 6동이 "
            "충돌하여 저·고 시나리오로 분리했다.",
            "빗물저장시설 2기의 배수는 General Summary를 근거로 한 "
            "시나리오이며 설계자 확인이 필요하다.",
            "수량 0은 실제 무공사보다 미완성 수식·설계값일 가능성이 있어 "
            "검토가 필요하다.",
            "공종별 서로 다른 단위의 수량은 합산해 비교하지 않는다.",
        ],
    }
    return {"summary": summary, "records": records}


def write_outputs(payload: dict[str, Any], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "KOICA_우간다_BOQ_구조화.json"
    csv_path = output_dir / "KOICA_우간다_BOQ_항목.csv"
    summary_path = output_dir / "KOICA_우간다_BOQ_추출요약.json"
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    summary_path.write_text(
        json.dumps(payload["summary"], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    records = payload["records"]
    if records:
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(records[0]))
            writer.writeheader()
            writer.writerows(records)


def compare_versions(
    primary_source: Path,
    primary: dict[str, Any],
    comparison_source: Path,
    comparison: dict[str, Any],
) -> dict[str, Any]:
    def keyed(payload: dict[str, Any]) -> dict[tuple[str, int], dict[str, Any]]:
        return {
            (row["source_sheet"], row["source_row"]): row
            for row in payload["records"]
        }

    primary_rows = keyed(primary)
    comparison_rows = keyed(comparison)
    shared = primary_rows.keys() & comparison_rows.keys()
    changed_descriptions: list[dict[str, Any]] = []
    item_quantity_changes: list[dict[str, Any]] = []
    for key in sorted(shared):
        current = primary_rows[key]
        previous = comparison_rows[key]
        if current["description"] != previous["description"]:
            changed_descriptions.append(
                {
                    "sheet": key[0],
                    "row": key[1],
                    "comparison": previous["description"],
                    "primary": current["description"],
                }
            )
        quantity_fields = (
            current["unit_raw"],
            current["quantity_base"],
            current["quantity_formula"],
        )
        previous_quantity_fields = (
            previous["unit_raw"],
            previous["quantity_base"],
            previous["quantity_formula"],
        )
        if (
            current["row_type"] in {"item", "item_unquantified"}
            or previous["row_type"] in {"item", "item_unquantified"}
        ) and quantity_fields != previous_quantity_fields:
            item_quantity_changes.append(
                {
                    "sheet": key[0],
                    "row": key[1],
                    "comparison": {
                        "unit": previous["unit_raw"],
                        "quantity": previous["quantity_base"],
                        "formula": previous["quantity_formula"],
                    },
                    "primary": {
                        "unit": current["unit_raw"],
                        "quantity": current["quantity_base"],
                        "formula": current["quantity_formula"],
                    },
                }
            )
    return {
        "primary_bid_no": infer_bid_no(primary_source),
        "primary_source": str(primary_source),
        "primary_sha256": sha256(primary_source),
        "comparison_bid_no": infer_bid_no(comparison_source),
        "comparison_source": str(comparison_source),
        "comparison_sha256": sha256(comparison_source),
        "primary_rows": len(primary_rows),
        "comparison_rows": len(comparison_rows),
        "shared_rows": len(shared),
        "only_in_primary": len(primary_rows.keys() - comparison_rows.keys()),
        "only_in_comparison": len(comparison_rows.keys() - primary_rows.keys()),
        "description_rows_changed": len(changed_descriptions),
        "item_quantity_rows_changed": len(item_quantity_changes),
        "description_change_examples": changed_descriptions[:20],
        "item_quantity_change_examples": item_quantity_changes[:20],
    }


def default_primary_source() -> Path:
    latest_root = Path("data/unpacked/L2023-00009-1")
    matches = sorted(
        latest_root.rglob("Section 4. Bill of Quantities_unpriced.xlsx")
    )
    return matches[0] if matches else DEFAULT_SOURCE


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path)
    parser.add_argument("--comparison-source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("outputs/koica-only-pilot"),
    )
    args = parser.parse_args()
    source = args.source or default_primary_source()
    payload = extract_boq(source)
    if args.comparison_source and args.comparison_source.resolve() != source.resolve():
        comparison = extract_boq(args.comparison_source)
        payload["summary"]["version_comparison"] = compare_versions(
            source,
            payload,
            args.comparison_source,
            comparison,
        )
    write_outputs(payload, args.output_dir)
    print(
        json.dumps(
            {
                "output_dir": str(args.output_dir),
                "item_rows": payload["summary"]["item_rows"],
                "quantified_items": payload["summary"]["quantified_items"],
                "priced_items": payload["summary"]["priced_items"],
                "detail_sheets": payload["summary"]["detail_sheets"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Recover cached IKHCC design-estimate rates from hidden KOICA BOQ sheets."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PRIMARY_BID_NO = "L2025-00085-1"
RELATED_BID_NOS = ("L2025-00070-1", "L2025-00071-1", PRIMARY_BID_NO)
PROJECT_NO = "2020-00119"
FACILITY_FILES = (
    (
        "병원",
        "HOSPITAL",
        "BILL No1 -Architectural Works-01. HOSPITAL(BLANK BILL).xlsx",
    ),
    (
        "기계·전기동",
        "MEP BUILDING",
        "BILL No1 -Architectural Works-02. MEP BUILDING(BLANK BILL).xlsx",
    ),
    (
        "연결교량",
        "BRIDGE",
        "BILL No1 -Architectural Works-03. BRIDGE(BLANK BILL).xlsx",
    ),
)
VISIBLE_SHEET = "Works particular bill"
ASSEMBLY_SHEET = "일위대가목록"
COMPARISON_SHEET = "단가대비표"
SETTINGS_SHEET = " 공사설정 "


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def number(value: Any) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return 0.0
    return float(value)


def parse_amount(value: Any) -> float:
    match = re.search(r"\d[\d,]*(?:\.\d+)?", str(value or ""))
    return float(match.group(0).replace(",", "")) if match else 0.0


def normalize_work_code(value: Any) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if float(value).is_integer():
            return str(int(value)).zfill(6)
    text = str(value or "").strip()
    return text.zfill(6) if text.isdigit() and len(text) < 6 else text


def find_one(root: Path, filename: str) -> Path:
    matches = list(root.rglob(filename))
    if len(matches) != 1:
        raise RuntimeError(f"expected one {filename!r} under {root}, found {len(matches)}")
    return matches[0]


def relative(path: Path, project_root: Path) -> str:
    return str(path.resolve().relative_to(project_root.resolve()))


def load_rate_library(
    workbook_path: Path,
    project_root: Path,
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]], dict[str, Any]]:
    values = load_workbook(
        workbook_path,
        data_only=True,
        read_only=False,
        keep_links=False,
    )
    formulas = load_workbook(
        workbook_path,
        data_only=False,
        read_only=False,
        keep_links=False,
    )
    rows: list[dict[str, Any]] = []
    by_code: dict[str, dict[str, Any]] = {}
    source_file = relative(workbook_path, project_root)
    source_hash = sha256(workbook_path)

    assembly_values = values[ASSEMBLY_SHEET]
    assembly_formulas = formulas[ASSEMBLY_SHEET]
    for row_no in range(4, assembly_values.max_row + 1):
        code_value = assembly_values.cell(row_no, 1).value
        if not code_value:
            continue
        code = str(code_value)
        material = number(assembly_values.cell(row_no, 5).value)
        labor = number(assembly_values.cell(row_no, 6).value)
        expense = number(assembly_values.cell(row_no, 7).value)
        total = number(assembly_values.cell(row_no, 8).value)
        if total <= 0:
            total = material + labor + expense
        record = {
            "rate_type": "복합일위대가",
            "item_code": code,
            "description": str(assembly_values.cell(row_no, 2).value or ""),
            "specification": str(assembly_values.cell(row_no, 3).value or ""),
            "unit": str(assembly_values.cell(row_no, 4).value or ""),
            "material_rate_usd": material,
            "labor_rate_usd": labor,
            "expense_rate_usd": expense,
            "total_rate_usd": total,
            "source_price_candidates": "",
            "source_pages": "",
            "source_note": str(assembly_values.cell(row_no, 10).value or ""),
            "material_formula": str(assembly_formulas.cell(row_no, 5).value or ""),
            "labor_formula": str(assembly_formulas.cell(row_no, 6).value or ""),
            "expense_formula": str(assembly_formulas.cell(row_no, 7).value or ""),
            "total_formula": str(assembly_formulas.cell(row_no, 8).value or ""),
            "source_sheet": ASSEMBLY_SHEET,
            "source_row": row_no,
            "source_file": source_file,
            "source_sha256": source_hash,
        }
        rows.append(record)
        by_code[code] = record

    comparison_values = values[COMPARISON_SHEET]
    comparison_formulas = formulas[COMPARISON_SHEET]
    for row_no in range(5, comparison_values.max_row + 1):
        code_value = comparison_values.cell(row_no, 1).value
        if not code_value:
            continue
        code = str(code_value)
        material = number(comparison_values.cell(row_no, 15).value)
        labor = number(comparison_values.cell(row_no, 16).value)
        expense = number(comparison_values.cell(row_no, 22).value)
        total = material + labor + expense
        price_candidates = [
            comparison_values.cell(row_no, column).value
            for column in (5, 7, 9, 11, 13)
        ]
        pages = [
            comparison_values.cell(row_no, column).value
            for column in (6, 8, 10, 12, 14)
        ]
        record = {
            "rate_type": "직접자원단가",
            "item_code": code,
            "description": str(comparison_values.cell(row_no, 2).value or ""),
            "specification": str(comparison_values.cell(row_no, 3).value or ""),
            "unit": str(comparison_values.cell(row_no, 4).value or ""),
            "material_rate_usd": material,
            "labor_rate_usd": labor,
            "expense_rate_usd": expense,
            "total_rate_usd": total,
            "source_price_candidates": " | ".join(
                str(value) for value in price_candidates if value not in (None, "", 0)
            ),
            "source_pages": " | ".join(
                str(value) for value in pages if value not in (None, "")
            ),
            "source_note": str(comparison_values.cell(row_no, 24).value or ""),
            "material_formula": str(comparison_formulas.cell(row_no, 15).value or ""),
            "labor_formula": str(comparison_formulas.cell(row_no, 16).value or ""),
            "expense_formula": str(comparison_formulas.cell(row_no, 22).value or ""),
            "total_formula": "=O{0}+P{0}+V{0}".format(row_no),
            "source_sheet": COMPARISON_SHEET,
            "source_row": row_no,
            "source_file": source_file,
            "source_sha256": source_hash,
        }
        rows.append(record)
        by_code[code] = record

    settings_values = values[SETTINGS_SHEET]
    settings_fx = 0.0
    for row_no in range(1, settings_values.max_row + 1):
        if str(settings_values.cell(row_no, 1).value or "").strip() == "환율":
            settings_fx = number(settings_values.cell(row_no, 2).value)
            break
    metadata = {
        "comparison_exchange_rate_krw_per_usd": number(comparison_values["AD3"].value),
        "comparison_markup_factor": number(comparison_values["AE3"].value),
        "settings_exchange_rate_krw_per_usd": settings_fx,
        "assembly_sheet_state": assembly_values.sheet_state,
        "comparison_sheet_state": comparison_values.sheet_state,
        "settings_sheet_state": settings_values.sheet_state,
    }
    values.close()
    formulas.close()
    return rows, by_code, metadata


def extract_boq_items(
    workbook_path: Path,
    facility_ko: str,
    facility_en: str,
    rates_by_code: dict[str, dict[str, Any]],
    project_root: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    values = load_workbook(
        workbook_path,
        data_only=True,
        read_only=False,
        keep_links=False,
    )
    formulas = load_workbook(
        workbook_path,
        data_only=False,
        read_only=False,
        keep_links=False,
    )
    value_sheet = values[VISIBLE_SHEET]
    formula_sheet = formulas[VISIBLE_SHEET]
    source_file = relative(workbook_path, project_root)
    source_hash = sha256(workbook_path)
    items: list[dict[str, Any]] = []
    visible_component_prices_blank = 0
    current_section_code = ""
    current_section_name = ""
    for row_no in range(4, value_sheet.max_row + 1):
        quantity_value = value_sheet.cell(row_no, 4).value
        code_value = value_sheet.cell(row_no, 14).value
        description_value = value_sheet.cell(row_no, 1).value
        description_text = str(description_value or "").strip()
        heading = re.match(r"^(\d{6,8})\s+.+", description_text)
        if heading and not isinstance(quantity_value, (int, float)):
            current_section_code = heading.group(1)
            current_section_name = description_text
        if (
            not isinstance(quantity_value, (int, float))
            or not code_value
            or not description_value
        ):
            continue
        code = str(code_value)
        rate = rates_by_code.get(code)
        material_rate = number(rate.get("material_rate_usd")) if rate else 0.0
        labor_rate = number(rate.get("labor_rate_usd")) if rate else 0.0
        expense_rate = number(rate.get("expense_rate_usd")) if rate else 0.0
        total_rate = number(rate.get("total_rate_usd")) if rate else 0.0
        quantity = float(quantity_value)
        visible_prices = [
            formula_sheet.cell(row_no, column).value for column in (5, 7, 9)
        ]
        if all(value in (None, "") for value in visible_prices):
            visible_component_prices_blank += 1
        work_code = current_section_code or normalize_work_code(
            value_sheet.cell(row_no, 17).value
        )
        record = {
            "facility_ko": facility_ko,
            "facility_en": facility_en,
            "section_code": work_code,
            "section_name": current_section_name or work_code,
            "description": str(description_value),
            "specification": str(value_sheet.cell(row_no, 2).value or ""),
            "unit": str(value_sheet.cell(row_no, 3).value or ""),
            "quantity": quantity,
            "item_code": code,
            "rate_source_type": rate.get("rate_type", "") if rate else "",
            "recovered_material_rate_usd": material_rate,
            "recovered_labor_rate_usd": labor_rate,
            "recovered_expense_rate_usd": expense_rate,
            "recovered_total_rate_usd": total_rate,
            "recovered_material_amount_usd": quantity * material_rate,
            "recovered_labor_amount_usd": quantity * labor_rate,
            "recovered_expense_amount_usd": quantity * expense_rate,
            "recovered_total_amount_usd": quantity * total_rate,
            "join_status": "복원" if rate and total_rate > 0 else "미복원",
            "visible_material_price": formula_sheet.cell(row_no, 5).value,
            "visible_labor_price": formula_sheet.cell(row_no, 7).value,
            "visible_expense_price": formula_sheet.cell(row_no, 9).value,
            "visible_total_price_formula": str(
                formula_sheet.cell(row_no, 11).value or ""
            ),
            "visible_total_price_cached": number(value_sheet.cell(row_no, 11).value),
            "source_sheet": VISIBLE_SHEET,
            "source_row": row_no,
            "source_file": source_file,
            "source_sha256": source_hash,
        }
        items.append(record)
    values.close()
    formulas.close()
    return items, {
        "facility_ko": facility_ko,
        "facility_en": facility_en,
        "item_rows": len(items),
        "unique_item_codes": len({row["item_code"] for row in items}),
        "recovered_rows": sum(row["join_status"] == "복원" for row in items),
        "visible_component_prices_blank_rows": visible_component_prices_blank,
        "recovered_total_amount_usd": sum(
            row["recovered_total_amount_usd"] for row in items
        ),
        "source_file": source_file,
        "source_sha256": source_hash,
    }


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise RuntimeError(f"cannot write empty CSV: {path}")
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def extract(project_root: Path) -> dict[str, Any]:
    project_root = project_root.resolve()
    primary_root = project_root / "data" / "unpacked" / PRIMARY_BID_NO
    workbook_paths = {
        facility_en: find_one(primary_root, filename)
        for _, facility_en, filename in FACILITY_FILES
    }
    hospital_path = workbook_paths["HOSPITAL"]
    rate_library, rates_by_code, rate_metadata = load_rate_library(
        hospital_path,
        project_root,
    )
    if len(rate_library) != len(rates_by_code):
        raise RuntimeError("hidden rate library contains duplicate item codes")

    items: list[dict[str, Any]] = []
    facility_summaries: list[dict[str, Any]] = []
    library_checks: list[dict[str, Any]] = []
    baseline_signature = {
        code: (
            row["material_rate_usd"],
            row["labor_rate_usd"],
            row["expense_rate_usd"],
            row["total_rate_usd"],
        )
        for code, row in rates_by_code.items()
    }
    for facility_ko, facility_en, _ in FACILITY_FILES:
        path = workbook_paths[facility_en]
        facility_rates, facility_rate_map, _ = load_rate_library(path, project_root)
        facility_signature = {
            code: (
                row["material_rate_usd"],
                row["labor_rate_usd"],
                row["expense_rate_usd"],
                row["total_rate_usd"],
            )
            for code, row in facility_rate_map.items()
        }
        library_checks.append(
            {
                "facility_en": facility_en,
                "rate_rows": len(facility_rates),
                "same_as_hospital": facility_signature == baseline_signature,
            }
        )
        facility_items, facility_summary = extract_boq_items(
            path,
            facility_ko,
            facility_en,
            rates_by_code,
            project_root,
        )
        items.extend(facility_items)
        facility_summaries.append(facility_summary)

    details = json.loads(
        (project_root / "data" / "manifests" / "details.json").read_text(
            encoding="utf-8"
        )
    )[PRIMARY_BID_NO]
    ceiling_usd = parse_amount(details["fields"].get("집행한도금액(달러)"))
    official_candidate_path = (
        project_root
        / "outputs"
        / "koica-official-open-data"
        / "KOICA_공식데이터_건축관련후보.csv"
    )
    official_record: dict[str, str] = {}
    with official_candidate_path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if (
                row.get("project_no_linked") == PROJECT_NO
                and row.get("source_dataset") == "annual_procurement_plan"
            ):
                official_record = row
                break
    gross_floor_area_m2 = parse_amount(official_record.get("gross_floor_area_m2"))
    planned_ceiling_krw = parse_amount(official_record.get("amount_value"))

    total_summary = find_one(primary_root, "Bill _Total Cost Summary-Blank 250306.xlsx")
    total_book = load_workbook(
        total_summary,
        data_only=False,
        read_only=False,
        keep_links=False,
    )
    currency_labels = []
    for worksheet in total_book.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = str(cell.value or "")
                if "USD" in value.upper():
                    currency_labels.append(
                        {
                            "sheet": worksheet.title,
                            "cell": cell.coordinate,
                            "value": value,
                        }
                    )
    total_book.close()

    version_lineage = []
    for bid_no in RELATED_BID_NOS:
        bid_root = project_root / "data" / "unpacked" / bid_no
        for _, facility_en, filename in FACILITY_FILES:
            path = find_one(bid_root, filename)
            version_lineage.append(
                {
                    "bid_no": bid_no,
                    "facility_en": facility_en,
                    "source_file": relative(path, project_root),
                    "sha256": sha256(path),
                }
            )

    facility_totals = {
        row["facility_ko"]: row["recovered_total_amount_usd"]
        for row in facility_summaries
    }
    total_recovered = sum(facility_totals.values())
    section_groups: dict[tuple[str, str, str], dict[str, Any]] = {}
    for item in items:
        key = (item["facility_ko"], item["section_code"], item["section_name"])
        group = section_groups.setdefault(
            key,
            {
                "facility_ko": item["facility_ko"],
                "section_code": item["section_code"],
                "section_name": item["section_name"],
                "item_rows": 0,
                "recovered_rows": 0,
                "material_amount_usd": 0.0,
                "labor_amount_usd": 0.0,
                "expense_amount_usd": 0.0,
                "total_amount_usd": 0.0,
            },
        )
        group["item_rows"] += 1
        group["recovered_rows"] += item["join_status"] == "복원"
        group["material_amount_usd"] += item["recovered_material_amount_usd"]
        group["labor_amount_usd"] += item["recovered_labor_amount_usd"]
        group["expense_amount_usd"] += item["recovered_expense_amount_usd"]
        group["total_amount_usd"] += item["recovered_total_amount_usd"]
    section_rows = sorted(
        section_groups.values(),
        key=lambda row: (row["facility_ko"], row["section_code"]),
    )

    output = project_root / "outputs" / "koica-only-pilot"
    write_csv(output / "KOICA_IKHCC_숨김설계단가_라이브러리.csv", rate_library)
    write_csv(output / "KOICA_IKHCC_BOQ_수량단가결합.csv", items)
    write_csv(output / "KOICA_IKHCC_공종요약.csv", section_rows)
    summary = {
        "project_no": PROJECT_NO,
        "project_name": details["fields"].get("사업명", ""),
        "representative_bid_no": PRIMARY_BID_NO,
        "related_bid_nos": list(RELATED_BID_NOS),
        "bid_title": details["title"],
        "notice_date": details["notice_date"],
        "detail_url": details["detail_url"],
        "source_boundary": "KOICA 현지입찰 공개 첨부의 숨김·캐시 설계견적 단가",
        "currency": "USD",
        "currency_evidence": {
            "source_file": relative(total_summary, project_root),
            "sha256": sha256(total_summary),
            "labels": currency_labels,
        },
        "rate_library_rows": len(rate_library),
        "assembly_rate_rows": sum(
            row["rate_type"] == "복합일위대가" for row in rate_library
        ),
        "direct_resource_rate_rows": sum(
            row["rate_type"] == "직접자원단가" for row in rate_library
        ),
        "positive_rate_rows": sum(row["total_rate_usd"] > 0 for row in rate_library),
        "boq_item_rows": len(items),
        "unique_boq_item_codes": len({row["item_code"] for row in items}),
        "recovered_item_rows": sum(row["join_status"] == "복원" for row in items),
        "recovery_rate": (
            sum(row["join_status"] == "복원" for row in items) / len(items)
            if items
            else 0
        ),
        "visible_component_prices_blank_rows": sum(
            row["visible_component_prices_blank_rows"] for row in facility_summaries
        ),
        "facility_summaries": facility_summaries,
        "library_consistency_checks": library_checks,
        "rate_metadata": rate_metadata,
        "recovered_architectural_bill_usd": total_recovered,
        "local_bid_ceiling_usd": ceiling_usd,
        "recovered_share_of_bid_ceiling": (
            total_recovered / ceiling_usd if ceiling_usd else None
        ),
        "gross_floor_area_m2": gross_floor_area_m2,
        "recovered_architectural_bill_usd_per_gfa_m2": (
            total_recovered / gross_floor_area_m2 if gross_floor_area_m2 else None
        ),
        "local_bid_ceiling_usd_per_gfa_m2": (
            ceiling_usd / gross_floor_area_m2 if gross_floor_area_m2 else None
        ),
        "planned_bid_ceiling_krw": planned_ceiling_krw,
        "planned_bid_ceiling_krw_per_gfa_m2": (
            planned_ceiling_krw / gross_floor_area_m2
            if gross_floor_area_m2
            else None
        ),
        "official_plan_source_url": official_record.get("official_record_url", ""),
        "version_lineage": version_lineage,
        "interpretation_limits": [
            "복원 단가는 입찰자 제시가·계약가·변경계약가·준공가가 아니라 설계견적 기초단가다.",
            "보이는 BOQ 단가 3개 열은 561개 항목 모두 공란이고, 숨김 시트의 캐시값을 품목코드로 결합했다.",
            "이번 합계는 Bill No.01 건축공사(병원·MEP동·교량)만 포함하며 토목·기계·전기·통신·소방·동원·사인 공사는 제외한다.",
            "단가대비표 환율 1,350.3과 공사설정 환율 1,289가 불일치하므로 적용 시점과 최종 환율을 재확인해야 한다.",
            "공고 3건의 대표 워크북 SHA-256은 동일하므로 독립 비용표본 3건으로 세지 않는다.",
            "연면적 5,486㎡와 계획 입찰한도액은 KOICA 연간발주계획의 계획값이며 변경 가능하다.",
        ],
    }
    (output / "KOICA_IKHCC_숨김단가_복원요약.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output / "KOICA_IKHCC_숨김단가_구조화.json").write_text(
        json.dumps(
            {
                "summary": summary,
                "rate_library": rate_library,
                "boq_items": items,
                "section_rows": section_rows,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return summary


def main() -> int:
    project_root = Path(__file__).resolve().parents[1]
    print(json.dumps(extract(project_root), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
